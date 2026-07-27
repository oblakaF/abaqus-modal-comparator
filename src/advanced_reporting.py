from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.figure import Figure
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font
from pypdf import PdfReader, PdfWriter

import reporting
from advanced_metrics import advanced_metric_summary, render_automac_comac


_INSTALLED = False
_ORIGINAL_EXPORT_EXCEL = None
_ORIGINAL_EXPORT_PDF = None


def export_excel_advanced(result, output_path: Path, image_directory: Path) -> Path:
    output_path = _ORIGINAL_EXPORT_EXCEL(result, output_path, image_directory)
    image_directory = Path(image_directory)
    image_path = render_automac_comac(
        result, image_directory / "automac_comac.png"
    )
    summary_values = advanced_metric_summary(result)

    workbook = load_workbook(output_path)
    if "AutoMAC COMAC" in workbook.sheetnames:
        del workbook["AutoMAC COMAC"]
    sheet = workbook.create_sheet("AutoMAC COMAC")
    sheet["A1"] = "AutoMAC and COMAC diagnostics"
    sheet["A1"].font = Font(size=16, bold=True)
    row = 3
    for key, value in summary_values.items():
        sheet.cell(row=row, column=1, value=key)
        sheet.cell(row=row, column=2, value=value)
        row += 1
    sheet.add_image(ExcelImage(str(image_path)), "A9")
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 20
    workbook.save(output_path)
    return output_path


def export_pdf_advanced(result, output_path: Path, image_directory: Path) -> Path:
    output_path = Path(output_path)
    image_directory = Path(image_directory)
    image_directory.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory() as temporary_directory:
        temporary = Path(temporary_directory)
        base_pdf = temporary / "base.pdf"
        extension_pdf = temporary / "advanced.pdf"
        _ORIGINAL_EXPORT_PDF(result, base_pdf, image_directory)
        image_path = render_automac_comac(
            result, image_directory / "automac_comac.png"
        )
        summary_values = advanced_metric_summary(result)

        with PdfPages(extension_pdf) as pdf:
            figure = Figure(figsize=(11.69, 8.27), constrained_layout=True)
            axis = figure.add_subplot(111)
            axis.imshow(__import__("matplotlib.image").image.imread(image_path))
            axis.set_title("AutoMAC and COMAC diagnostics", fontsize=16)
            axis.axis("off")
            pdf.savefig(figure)

            figure = Figure(figsize=(11.69, 8.27))
            figure.text(0.06, 0.93, "Advanced modal diagnostics", fontsize=20, weight="bold")
            y = 0.85
            for key, value in summary_values.items():
                figure.text(0.08, y, f"{key}: {value:.4f}", fontsize=12)
                y -= 0.06
            figure.text(
                0.08,
                y - 0.02,
                "High off-diagonal AutoMAC indicates spatial aliasing or modes that the measurement grid cannot distinguish.\n"
                "Low COMAC identifies measurement locations where Abaqus and experiment disagree consistently.",
                fontsize=11,
            )
            pdf.savefig(figure)

        writer = PdfWriter()
        for source in (base_pdf, extension_pdf):
            reader = PdfReader(str(source))
            for page in reader.pages:
                writer.add_page(page)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("wb") as stream:
            writer.write(stream)
    return output_path


def install_advanced_reporting() -> None:
    global _INSTALLED, _ORIGINAL_EXPORT_EXCEL, _ORIGINAL_EXPORT_PDF
    if _INSTALLED:
        return
    _ORIGINAL_EXPORT_EXCEL = reporting.export_excel
    _ORIGINAL_EXPORT_PDF = reporting.export_pdf
    reporting.export_excel = export_excel_advanced
    reporting.export_pdf = export_pdf_advanced
    _INSTALLED = True
