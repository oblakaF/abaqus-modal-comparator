from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Dict, List

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader, PdfWriter

import reporting
from modal_core import ComparisonResult


_ORIGINAL_EXPORT_EXCEL = None
_ORIGINAL_EXPORT_PDF = None
_INSTALLED = False


def _metadata_list(result: ComparisonResult, key: str) -> List[Dict]:
    value = result.abaqus.metadata.get(key, [])
    return value if isinstance(value, list) else []


def render_frf_diagnostics(result: ComparisonResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    metadata = result.experimental.metadata
    frequency = np.asarray(metadata.get("_frf_frequency_hz", []), dtype=float)
    indicator = np.asarray(metadata.get("_frf_indicator", []), dtype=float)
    coherence = np.asarray(metadata.get("_frf_mean_coherence", []), dtype=float)

    figure, primary = plt.subplots(figsize=(10.5, 5.8), constrained_layout=True)
    if len(frequency) == 0 or len(indicator) != len(frequency):
        primary.text(
            0.5,
            0.5,
            "FRF diagnostic data are not available for this experimental file.",
            ha="center",
            va="center",
            transform=primary.transAxes,
        )
        primary.axis("off")
        figure.savefig(output_path, dpi=170)
        plt.close(figure)
        return output_path

    normalized = indicator / max(float(np.max(indicator)), 1e-30)
    primary.semilogy(frequency, np.maximum(normalized, 1e-12), label="Combined FRF indicator")
    primary.set_xlabel("Frequency, Hz")
    primary.set_ylabel("Normalized FRF indicator")
    modal_set_name = metadata.get("modal_set_name")
    if metadata.get("dataset_58_role") == "diagnostic_only" and modal_set_name:
        primary.set_title(
            "Dataset 58 FRF diagnostics; fitted modes from dataset 55 / "
            f"{modal_set_name}"
        )
    primary.grid(True, alpha=0.25)

    all_peaks = [mode.frequency_hz for mode in result.experimental.sorted_modes()]
    matched_experimental = {pair.experimental_mode: pair for pair in result.pairs}
    for mode in result.experimental.sorted_modes():
        pair = matched_experimental.get(mode.number)
        if pair is None:
            primary.axvline(mode.frequency_hz, linewidth=0.7, alpha=0.22)
        else:
            primary.axvline(mode.frequency_hz, linewidth=1.5, alpha=0.85)
            primary.text(
                mode.frequency_hz,
                0.92,
                f"E{mode.number}",
                rotation=90,
                va="top",
                ha="right",
                transform=primary.get_xaxis_transform(),
                fontsize=8,
            )
            primary.axvline(pair.abaqus_frequency_hz, linestyle="--", linewidth=1.0, alpha=0.7)

    # Legacy metadata (saved before coherence_status existed) is normalized to
    # "unavailable" instead of silently skipping the warning below.
    coherence_status = str(metadata.get("coherence_status") or "unavailable")
    if len(coherence) == len(frequency) and coherence_status == "computed":
        secondary = primary.twinx()
        secondary.plot(frequency, np.clip(coherence, 0.0, 1.0), alpha=0.45, label="Mean coherence")
        secondary.set_ylim(0.0, 1.05)
        secondary.set_ylabel("Mean coherence")
    else:
        primary.text(
            0.02,
            0.02,
            "Coherence: "
            + ("not available (no dataset-58 coherence channels)" if coherence_status == "unavailable"
               else "could not be parsed — see warnings"),
            transform=primary.transAxes,
            fontsize=8,
            color="firebrick",
        )

    primary.set_title(
        "Experimental FRF diagnostics\n"
        "Solid vertical lines: matched experimental peaks; dashed lines: matched Abaqus frequencies"
    )
    figure.savefig(output_path, dpi=170)
    plt.close(figure)
    return output_path


def render_verified_mac_matrix(result: ComparisonResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    abaqus_lookup = {mode: index for index, mode in enumerate(result.abaqus_mode_numbers)}
    experiment_lookup = {
        mode: index for index, mode in enumerate(result.experimental_mode_numbers)
    }
    matrix = np.full((len(result.pairs), len(result.pairs)), np.nan, dtype=float)
    row_labels = []
    column_labels = []
    for row, pair_row in enumerate(result.pairs):
        row_labels.append(f"A{pair_row.abaqus_mode}")
        for column, pair_column in enumerate(result.pairs):
            if row == 0:
                column_labels.append(f"E{pair_column.experimental_mode}")
            matrix[row, column] = result.mac_matrix[
                abaqus_lookup[pair_row.abaqus_mode],
                experiment_lookup[pair_column.experimental_mode],
            ]

    figure, axis = plt.subplots(figsize=(7.2, 5.8), constrained_layout=True)
    if not result.pairs:
        axis.text(
            0.5,
            0.5,
            "No accepted mode pairs\nSee the full MAC matrix and candidate diagnostics",
            ha="center",
            va="center",
            transform=axis.transAxes,
        )
        axis.set_axis_off()
        axis.set_title("Verified-pair MAC submatrix")
        figure.savefig(output_path, dpi=170)
        plt.close(figure)
        return output_path
    image = axis.imshow(matrix, vmin=0.0, vmax=1.0, aspect="auto")
    figure.colorbar(image, ax=axis, label="MAC")
    axis.set_xticks(range(len(column_labels)))
    axis.set_xticklabels(column_labels)
    axis.set_yticks(range(len(row_labels)))
    axis.set_yticklabels(row_labels)
    axis.set_xlabel("Verified experimental modes")
    axis.set_ylabel("Verified Abaqus modes")
    axis.set_title("Verified-pair MAC submatrix")
    for row in range(matrix.shape[0]):
        for column in range(matrix.shape[1]):
            value = matrix[row, column]
            if np.isfinite(value):
                axis.text(column, row, f"{value:.2f}", ha="center", va="center", fontsize=9)
    figure.savefig(output_path, dpi=170)
    plt.close(figure)
    return output_path


def quality_control_text(result: ComparisonResult) -> str:
    metadata = result.abaqus.metadata
    quality = metadata.get("quality_control", {})
    excluded = quality.get("excluded_abaqus_modes", [])
    unmatched = _metadata_list(result, "unmatched_abaqus_modes")
    experimental = _metadata_list(result, "unmatched_experimental_candidates")
    close_groups = _metadata_list(result, "close_abaqus_mode_groups")

    lines = [
        "QUALITY CONTROL SUMMARY",
        "=" * 72,
        f"Reliable matched pairs: {len(result.pairs)}",
        f"Rigid / near-zero Abaqus modes excluded: {excluded or 'none'}",
    ]
    experimental_metadata = result.experimental.metadata
    if experimental_metadata.get("modal_set_key"):
        lines.extend(
            [
                "Experimental source: curve-fitted dataset 55",
                f"Experimental modal set: {experimental_metadata.get('modal_set_name')}",
                f"Experimental modal-set key: {experimental_metadata.get('modal_set_key')}",
                "Dataset-55 residual records excluded: "
                f"{experimental_metadata.get('excluded_residual_count', 0)}",
                f"Dataset-58 role: {experimental_metadata.get('dataset_58_role', 'unavailable')}",
            ]
        )
    if unmatched:
        lines.append(
            "Unmatched Abaqus modes: "
            + ", ".join(
                f"{item['mode']} ({item['frequency_hz']:.4g} Hz)" for item in unmatched
            )
        )
    else:
        lines.append("Unmatched Abaqus modes: none")

    if close_groups:
        lines.append("Close Abaqus mode groups:")
        for group in close_groups:
            lines.append(
                "  • "
                + " / ".join(
                    f"mode {item['mode']} ({item['frequency_hz']:.4g} Hz)"
                    for item in group
                )
                + " — possible mixed experimental shape"
            )
    else:
        lines.append("Close Abaqus mode groups: none")

    if experimental:
        lines.append(
            "Unmatched experimental peak candidates in the Abaqus frequency band: "
            + ", ".join(
                f"E{item['mode']} ({item['frequency_hz']:.4g} Hz)" for item in experimental
            )
        )

    lines.append("")
    lines.append("Accepted-pair limits:")
    lines.append(
        f"  frequency error <= {quality.get('maximum_accepted_frequency_error_percent', 15.0):.1f}%"
    )
    lines.append(f"  MAC >= {quality.get('minimum_accepted_mac', 0.50):.2f}")
    return "\n".join(lines)


def export_excel_enhanced(
    result: ComparisonResult,
    output_path: Path,
    image_directory: Path,
) -> Path:
    output_path = _ORIGINAL_EXPORT_EXCEL(result, output_path, image_directory)
    image_directory = Path(image_directory)
    frf_path = render_frf_diagnostics(result, image_directory / "frf_diagnostics.png")
    verified_path = render_verified_mac_matrix(
        result, image_directory / "verified_mac_matrix.png"
    )

    workbook = load_workbook(output_path)
    if "FRF Diagnostics" in workbook.sheetnames:
        del workbook["FRF Diagnostics"]
    sheet = workbook.create_sheet("FRF Diagnostics")
    sheet["A1"] = "Experimental FRF diagnostics"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet.add_image(ExcelImage(str(frf_path)), "A3")
    sheet.add_image(ExcelImage(str(verified_path)), "A34")

    if "Quality Control" in workbook.sheetnames:
        del workbook["Quality Control"]
    quality_sheet = workbook.create_sheet("Quality Control")
    quality_sheet["A1"] = "Quality-control decisions"
    quality_sheet["A1"].font = Font(size=16, bold=True)
    for row, line in enumerate(quality_control_text(result).splitlines(), start=3):
        quality_sheet.cell(row=row, column=1, value=line)
    quality_sheet.column_dimensions["A"].width = 120

    rejected = _metadata_list(result, "rejected_forced_pairs")
    if rejected:
        start = len(quality_control_text(result).splitlines()) + 6
        headers = [
            "Rejected Abaqus mode",
            "Forced experimental mode",
            "Abaqus frequency, Hz",
            "Experimental frequency, Hz",
            "Error, %",
            "MAC",
        ]
        for column, header in enumerate(headers, start=1):
            cell = quality_sheet.cell(row=start, column=column, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F4CCCC")
        for offset, item in enumerate(rejected, start=1):
            quality_sheet.append(
                [
                    item.get("abaqus_mode"),
                    item.get("experimental_mode"),
                    item.get("abaqus_frequency_hz"),
                    item.get("experimental_frequency_hz"),
                    item.get("frequency_error_percent"),
                    item.get("mac"),
                ]
            )

    workbook.save(output_path)
    return output_path


def export_pdf_enhanced(
    result: ComparisonResult,
    output_path: Path,
    image_directory: Path,
) -> Path:
    output_path = Path(output_path)
    image_directory = Path(image_directory)
    image_directory.mkdir(parents=True, exist_ok=True)

    with TemporaryDirectory() as temporary_directory:
        temporary = Path(temporary_directory)
        base_pdf = temporary / "base.pdf"
        extension_pdf = temporary / "extension.pdf"
        _ORIGINAL_EXPORT_PDF(result, base_pdf, image_directory)

        frf_path = render_frf_diagnostics(result, image_directory / "frf_diagnostics.png")
        verified_path = render_verified_mac_matrix(
            result, image_directory / "verified_mac_matrix.png"
        )
        with PdfPages(extension_pdf) as pdf:
            for image_path, title in (
                (frf_path, "FRF and coherence diagnostics"),
                (verified_path, "Verified-pair MAC submatrix"),
            ):
                figure, axis = plt.subplots(figsize=(11.69, 8.27))
                axis.imshow(plt.imread(image_path))
                axis.set_title(title, fontsize=16)
                axis.axis("off")
                pdf.savefig(figure)
                plt.close(figure)

            figure = plt.figure(figsize=(11.69, 8.27))
            figure.text(0.06, 0.94, "Quality-control summary", fontsize=20, weight="bold")
            y = 0.89
            for line in quality_control_text(result).splitlines():
                figure.text(0.07, y, line, fontsize=9, family="monospace")
                y -= 0.032
                if y < 0.05:
                    break
            plt.axis("off")
            pdf.savefig(figure)
            plt.close(figure)

        writer = PdfWriter()
        for source in (base_pdf, extension_pdf):
            reader = PdfReader(str(source))
            for page in reader.pages:
                writer.add_page(page)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("wb") as stream:
            writer.write(stream)
    return output_path


def install_reporting_enhancements() -> None:
    global _INSTALLED, _ORIGINAL_EXPORT_EXCEL, _ORIGINAL_EXPORT_PDF
    if _INSTALLED:
        return
    _ORIGINAL_EXPORT_EXCEL = reporting.export_excel
    _ORIGINAL_EXPORT_PDF = reporting.export_pdf
    reporting.export_excel = export_excel_enhanced
    reporting.export_pdf = export_pdf_enhanced
    _INSTALLED = True
