from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.tri as mtri
import numpy as np
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill

from modal_core import ComparisonResult, ModePairResult


def _plane_axes(coordinates: np.ndarray) -> Tuple[int, int, int]:
    spans = np.ptp(np.asarray(coordinates, dtype=float), axis=0)
    order = np.argsort(spans)
    normal_axis = int(order[0])
    plane_axes = [int(value) for value in order[1:]]
    return plane_axes[0], plane_axes[1], normal_axis


def _real_mode_component(coordinates: np.ndarray, vectors: np.ndarray) -> np.ndarray:
    vectors = np.asarray(vectors)
    if vectors.size == 0:
        return np.array([], dtype=float)

    flattened = vectors.reshape(-1)
    index = int(np.argmax(np.abs(flattened)))
    reference = flattened[index]
    if abs(reference) > 1e-30:
        vectors = vectors * np.exp(-1j * np.angle(reference))

    _, _, normal_axis = _plane_axes(coordinates)
    component = np.real(vectors[:, normal_axis])
    if np.max(np.abs(component)) <= 1e-12:
        component = np.real(vectors[:, np.argmax(np.linalg.norm(vectors, axis=0))])
    maximum = np.max(np.abs(component)) if len(component) else 0.0
    if maximum > 0:
        component = component / maximum
    return component


def render_mode_shape(
    coordinates: np.ndarray,
    vectors: np.ndarray,
    title: str,
    output_path: Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    coordinates = np.asarray(coordinates, dtype=float)
    component = _real_mode_component(coordinates, vectors)
    axis_x, axis_y, _ = _plane_axes(coordinates)
    x = coordinates[:, axis_x]
    y = coordinates[:, axis_y]

    figure, axis = plt.subplots(figsize=(6.0, 4.6), constrained_layout=True)
    plotted = False
    if len(x) >= 3 and len(np.unique(np.column_stack((x, y)), axis=0)) >= 3:
        try:
            triangulation = mtri.Triangulation(x, y)
            contour = axis.tricontourf(triangulation, component, levels=21)
            axis.tricontour(triangulation, component, levels=[0.0], linewidths=1.2)
            figure.colorbar(contour, ax=axis, label="Normalized modal amplitude")
            plotted = True
        except Exception:
            plotted = False

    if not plotted:
        scatter = axis.scatter(x, y, c=component, s=24)
        figure.colorbar(scatter, ax=axis, label="Normalized modal amplitude")

    axis.set_title(title)
    axis.set_xlabel(("X", "Y", "Z")[axis_x])
    axis.set_ylabel(("X", "Y", "Z")[axis_y])
    axis.set_aspect("equal", adjustable="box")
    axis.grid(True, alpha=0.25)
    figure.savefig(output_path, dpi=160)
    plt.close(figure)
    return output_path


def render_overlay(pair: ModePairResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    coordinates = np.asarray(pair.coordinates, dtype=float)
    abaqus_component = _real_mode_component(coordinates, pair.abaqus_vector)
    experimental_component = _real_mode_component(coordinates, pair.experimental_vector)
    axis_x, axis_y, _ = _plane_axes(coordinates)

    figure, axis = plt.subplots(figsize=(6.0, 4.6), constrained_layout=True)
    axis.scatter(
        coordinates[:, axis_x],
        coordinates[:, axis_y],
        c=experimental_component,
        marker="o",
        s=44,
        label="Experiment",
    )
    axis.scatter(
        coordinates[:, axis_x],
        coordinates[:, axis_y],
        c=abaqus_component,
        marker="x",
        s=34,
        label="Abaqus",
    )
    title = f"Mode {pair.abaqus_mode} / {pair.experimental_mode} overlay"
    title += f"\nMAC={pair.mac:.3f}" if pair.mac is not None else "\nMAC unavailable"
    axis.set_title(title)
    axis.set_xlabel(("X", "Y", "Z")[axis_x])
    axis.set_ylabel(("X", "Y", "Z")[axis_y])
    axis.set_aspect("equal", adjustable="box")
    axis.grid(True, alpha=0.25)
    axis.legend()
    figure.savefig(output_path, dpi=160)
    plt.close(figure)
    return output_path


def render_pair_images(pair: ModePairResult, output_directory: Path) -> Dict[str, Path]:
    output_directory = Path(output_directory)
    stem = f"abaqus_{pair.abaqus_mode}_experiment_{pair.experimental_mode}"
    return {
        "abaqus": render_mode_shape(
            pair.coordinates,
            pair.abaqus_vector,
            f"Abaqus mode {pair.abaqus_mode} — {pair.abaqus_frequency_hz:.3f} Hz",
            output_directory / f"{stem}_abaqus.png",
        ),
        "experimental": render_mode_shape(
            pair.coordinates,
            pair.experimental_vector,
            f"Experimental mode {pair.experimental_mode} — {pair.experimental_frequency_hz:.3f} Hz",
            output_directory / f"{stem}_experimental.png",
        ),
        "overlay": render_overlay(pair, output_directory / f"{stem}_overlay.png"),
    }


def render_mac_matrix(result: ComparisonResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    matrix = np.asarray(result.mac_matrix, dtype=float)
    figure, axis = plt.subplots(figsize=(8.0, 6.2), constrained_layout=True)
    image = axis.imshow(matrix, vmin=0.0, vmax=1.0, aspect="auto")
    figure.colorbar(image, ax=axis, label="MAC")
    axis.set_xticks(range(len(result.experimental_mode_numbers)))
    axis.set_xticklabels(result.experimental_mode_numbers)
    axis.set_yticks(range(len(result.abaqus_mode_numbers)))
    axis.set_yticklabels(result.abaqus_mode_numbers)
    axis.set_xlabel("Experimental mode")
    axis.set_ylabel("Abaqus mode")
    axis.set_title("Modal Assurance Criterion matrix")

    for row in range(matrix.shape[0]):
        for column in range(matrix.shape[1]):
            value = matrix[row, column]
            if np.isfinite(value):
                axis.text(column, row, f"{value:.2f}", ha="center", va="center", fontsize=8)

    figure.savefig(output_path, dpi=170)
    plt.close(figure)
    return output_path


def render_frequency_comparison(result: ComparisonResult, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    labels = [f"A{pair.abaqus_mode}/E{pair.experimental_mode}" for pair in result.pairs]
    abaqus_values = [pair.abaqus_frequency_hz for pair in result.pairs]
    experimental_values = [pair.experimental_frequency_hz for pair in result.pairs]
    positions = np.arange(len(labels))
    width = 0.38

    figure, axis = plt.subplots(figsize=(10.0, 5.0), constrained_layout=True)
    axis.bar(positions - width / 2, abaqus_values, width, label="Abaqus")
    axis.bar(positions + width / 2, experimental_values, width, label="Experiment")
    axis.set_xticks(positions)
    axis.set_xticklabels(labels, rotation=45, ha="right")
    axis.set_ylabel("Frequency, Hz")
    axis.set_title("Matched natural frequencies")
    axis.grid(True, axis="y", alpha=0.25)
    axis.legend()
    figure.savefig(output_path, dpi=170)
    plt.close(figure)
    return output_path


def export_excel(result: ComparisonResult, output_path: Path, image_directory: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image_directory = Path(image_directory)
    image_directory.mkdir(parents=True, exist_ok=True)

    mac_path = render_mac_matrix(result, image_directory / "mac_matrix.png")
    frequency_path = render_frequency_comparison(result, image_directory / "frequency_comparison.png")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Abaqus Modal Comparator"
    summary["A1"].font = Font(size=18, bold=True)
    summary["A3"] = "Abaqus source"
    summary["B3"] = str(result.abaqus.source_path)
    summary["A4"] = "Experimental source"
    summary["B4"] = str(result.experimental.source_path)
    summary["A5"] = "Matched mode pairs"
    summary["B5"] = len(result.pairs)
    summary["A6"] = "Geometry matched fraction"
    summary["B6"] = result.geometry.matched_fraction
    summary["B6"].number_format = "0.0%"
    summary["A7"] = "Geometry normalized RMS distance"
    summary["B7"] = result.geometry.normalized_rms_distance
    summary["B7"].number_format = "0.00%"
    summary["A8"] = "Detected coordinate scale"
    summary["B8"] = result.geometry.coordinate_scale

    if result.pairs:
        summary["A10"] = "Mean frequency error"
        summary["B10"] = float(np.mean([pair.frequency_error_percent for pair in result.pairs])) / 100.0
        summary["B10"].number_format = "0.00%"
        mac_values = [pair.mac for pair in result.pairs if pair.mac is not None]
        if mac_values:
            summary["A11"] = "Mean MAC"
            summary["B11"] = float(np.mean(mac_values))

    row = 13
    if result.warnings:
        summary.cell(row=row, column=1, value="Warnings").font = Font(bold=True)
        for warning in result.warnings:
            row += 1
            summary.cell(row=row, column=1, value=warning)

    summary.add_image(ExcelImage(str(frequency_path)), "D3")
    summary.add_image(ExcelImage(str(mac_path)), "D28")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 75

    comparison = workbook.create_sheet("Mode Comparison")
    headers = [
        "Abaqus mode",
        "Experimental mode",
        "Abaqus frequency, Hz",
        "Experimental frequency, Hz",
        "Frequency error, %",
        "MAC",
        "Order changed",
        "Mapped points",
        "Status",
    ]
    comparison.append(headers)
    for cell in comparison[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    for pair in result.pairs:
        comparison.append(
            [
                pair.abaqus_mode,
                pair.experimental_mode,
                pair.abaqus_frequency_hz,
                pair.experimental_frequency_hz,
                pair.frequency_error_percent,
                pair.mac,
                "Yes" if pair.order_changed else "No",
                pair.mapped_points,
                pair.status,
            ]
        )
    comparison.freeze_panes = "A2"
    comparison.auto_filter.ref = comparison.dimensions
    for column in range(1, len(headers) + 1):
        comparison.column_dimensions[chr(64 + column)].width = 22

    mac_sheet = workbook.create_sheet("MAC Matrix")
    mac_sheet.append(["Abaqus / Experiment"] + result.experimental_mode_numbers)
    for index, abaqus_mode in enumerate(result.abaqus_mode_numbers):
        row_values: List[object] = [abaqus_mode]
        for value in result.mac_matrix[index]:
            row_values.append(None if not np.isfinite(value) else float(value))
        mac_sheet.append(row_values)

    geometry = workbook.create_sheet("Geometry Match")
    geometry.append(["Experimental point", "Abaqus node index", "Distance"])
    for index, (mapped_index, distance) in enumerate(
        zip(result.geometry.experimental_to_abaqus, result.geometry.distances),
        start=1,
    ):
        geometry.append([index, int(mapped_index), float(distance)])

    history = workbook.create_sheet("Abaqus History")
    history.append(["Region", "Output", "Description", "X", "Value"])
    for item in result.abaqus.history:
        for x_value, y_value in item.get("data", []):
            history.append(
                [
                    item.get("region"),
                    item.get("name"),
                    item.get("description"),
                    x_value,
                    y_value,
                ]
            )

    workbook.save(output_path)
    return output_path


def export_pdf(result: ComparisonResult, output_path: Path, image_directory: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image_directory = Path(image_directory)
    image_directory.mkdir(parents=True, exist_ok=True)

    mac_path = render_mac_matrix(result, image_directory / "mac_matrix.png")
    frequency_path = render_frequency_comparison(result, image_directory / "frequency_comparison.png")

    with PdfPages(output_path) as pdf:
        figure = plt.figure(figsize=(11.69, 8.27))
        figure.text(0.07, 0.92, "Abaqus Modal Comparator", fontsize=22, weight="bold")
        figure.text(0.07, 0.86, f"Abaqus: {result.abaqus.source_path}", fontsize=10)
        figure.text(0.07, 0.82, f"Experiment: {result.experimental.source_path}", fontsize=10)
        figure.text(0.07, 0.76, f"Matched pairs: {len(result.pairs)}", fontsize=12)
        figure.text(0.07, 0.72, f"Geometry match: {result.geometry.matched_fraction:.1%}", fontsize=12)
        figure.text(0.07, 0.68, f"Geometry RMS / diagonal: {result.geometry.normalized_rms_distance:.2%}", fontsize=12)
        if result.warnings:
            figure.text(0.07, 0.60, "Warnings", fontsize=13, weight="bold")
            y = 0.56
            for warning in result.warnings:
                figure.text(0.09, y, "• " + warning, fontsize=10)
                y -= 0.04
        plt.axis("off")
        pdf.savefig(figure)
        plt.close(figure)

        for image_path, title in ((frequency_path, "Frequency comparison"), (mac_path, "MAC matrix")):
            image = plt.imread(image_path)
            figure, axis = plt.subplots(figsize=(11.69, 8.27))
            axis.imshow(image)
            axis.set_title(title, fontsize=16)
            axis.axis("off")
            pdf.savefig(figure)
            plt.close(figure)

        for pair in result.pairs:
            paths = render_pair_images(pair, image_directory / "mode_pairs")
            figure, axes = plt.subplots(1, 3, figsize=(11.69, 8.27), constrained_layout=True)
            for axis, key, title in zip(
                axes,
                ("abaqus", "experimental", "overlay"),
                ("Abaqus", "Experiment", "Overlay"),
            ):
                axis.imshow(plt.imread(paths[key]))
                axis.set_title(title)
                axis.axis("off")
            figure.suptitle(
                f"Abaqus mode {pair.abaqus_mode} ↔ experimental mode {pair.experimental_mode}\n"
                f"Frequency error {pair.frequency_error_percent:.2f}% | "
                + (f"MAC {pair.mac:.3f}" if pair.mac is not None else "MAC unavailable"),
                fontsize=15,
            )
            pdf.savefig(figure)
            plt.close(figure)

    return output_path
