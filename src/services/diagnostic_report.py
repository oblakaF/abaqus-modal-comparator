"""Diagnostic PDF / Excel report of an ODB + UNV analysis without a physical registration (CORE-E2E-D).

When the comparator cannot establish the physical registration (``FREQUENCY_ONLY_DIAGNOSTIC``, or an experimental file
without usable data) there is no ``ComparisonResult`` and the comparison report does not apply. This module writes a
self-contained engineering report of what the analysis actually produced, from the existing backend objects only:

- ``LoadedModalData`` (the datasets the run loaded, the registration state and its reason);
- ``FrequencyOnlyDiagnostic`` (``services.new_odb_comparison``: frequencies, nearest-frequency rows, range notes);
- the existing classifications of ``services.modal_data_viewer`` (FE mode, curve-fitted mode, FRF peak snapshot);
- the guided-registration evidence list (``services.guided_registration.REQUIRED_EVIDENCE``) for what is missing.

It never computes a MAC, never pairs modes, never chooses an orientation, never writes a pair frequency error and never
states a material property. Nearest frequency is not physical mode pairing.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
import os
from pathlib import Path
import tempfile
import textwrap
from typing import Any, Mapping, Optional, Sequence

from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.figure import Figure
import matplotlib.image as mpimg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scientific_state import modal_dataset_geometry_identity

from . import modal_data_viewer as viewer

TITLE = "Abaqus–Experiment Modal Diagnostic Report"
BANNER = "FREQUENCY-ONLY DIAGNOSTIC — PHYSICAL MODE PAIRING NOT ESTABLISHED"
FE_ONLY_BANNER = "DIAGNOSTIC — EXPERIMENTAL DATA NOT USABLE; NO FE–EXPERIMENT COMPARISON"
NOT_PAIRING = "Nearest frequency ≠ physical mode pairing."
NOT_COMPUTED = "NOT_COMPUTED"
SCHEMA = "modal-diagnostic-report/v1"
STATUS_TEXT = {viewer.FITTED_MODE: "FITTED MODE (curve-fitted modal set)",
               viewer.FRF_PEAK_SNAPSHOT: "PROVISIONAL — FRF peak, NOT a validated fitted mode",
               viewer.UNCLASSIFIED: "UNCLASSIFIED experimental record"}
LIMITS = (
    "No physical registration: no measured point is mapped to an FE node.",
    "No MAC, no mode pairs, no pair frequency error and no orientation were computed or chosen.",
    "Nearest-frequency rows describe frequency proximity only; they are not mode pairs and carry no shape evidence.",
    "No material property (E1, E2, G12, ...) is identified or implied; no Auto-ID observation is created.",
)


class ReportWriteError(Exception):
    """The report could not be written to the chosen file (open in another program, no permission, ...)."""


def _identity_of(path: Optional[Path], cancel=None) -> dict:
    if path is None:
        return {"path": None, "size": None, "sha256": None, "note": "not recorded"}
    from services.archived_sensitivity_evidence import AcceptedReferenceRefusal, fresh_content_identity

    try:
        identity = fresh_content_identity(Path(path), cancel)
    except AcceptedReferenceRefusal as error:  # unreadable or changed while hashed: stated, never invented
        return {"path": str(path), "size": None, "sha256": None, "note": f"not available: {error}"}
    return {"path": str(path), "size": identity["size"], "sha256": identity["sha256"], "note": ""}


@dataclass(frozen=True)
class DiagnosticReport:
    created: str
    status: str
    banner: str
    sources: tuple[tuple[str, str], ...]
    reason: str
    checks_passed: tuple[str, ...]
    checks_not_passed: tuple[str, ...]
    needed: tuple[str, ...]
    fe_rows: tuple[Mapping[str, Any], ...]
    fe_range_hz: Optional[tuple[float, float]]
    experimental_rows: tuple[Mapping[str, Any], ...]
    proximity_rows: tuple[Mapping[str, Any], ...]
    proximity_source: str
    conclusions: tuple[str, ...]
    provenance: Mapping[str, Any]
    images: tuple[tuple[str, str], ...] = field(default_factory=tuple)

    @property
    def mac(self) -> str:
        return NOT_COMPUTED


def _fe_rows(abaqus) -> tuple[list[dict], Optional[tuple[float, float]]]:
    rows = [{"mode": entry.number, "frequency_hz": entry.frequency_hz, "points": entry.point_count,
             "components": "/".join({"x": "U1", "y": "U2", "z": "U3"}[axis] for axis in entry.axes) or "none",
             "shape": "spatial view available" if entry.spatial_available else entry.spatial_reason}
            for entry in viewer.viewer_modes(abaqus, viewer.FE)]
    hz = [row["frequency_hz"] for row in rows]
    return rows, ((min(hz), max(hz)) if hz else None)


def _range_note(frequency: float, fe_range) -> str:
    if fe_range is None:
        return "no FE frequency extracted"
    low, high = fe_range
    if frequency < low:
        return f"below the lowest extracted FE mode ({low:.2f} Hz)"
    if frequency > high:
        return f"above the highest extracted FE mode ({high:.2f} Hz); extend the Abaqus mode range"
    return ""


def _experimental_rows(experimental, fe_range) -> list[dict]:
    rows = []
    modal_set = (experimental.metadata.get("modal_set_name") or experimental.metadata.get("modal_set_key") or
                 "none (no fitted modal set)") if experimental is not None else ""
    for entry in viewer.viewer_modes(experimental, viewer.EXPERIMENTAL):
        warnings = []
        note = _range_note(entry.frequency_hz, fe_range)
        if note:
            warnings.append(note)
        if not entry.spatial_available:
            warnings.append(entry.spatial_reason)
        quality = dict(entry.quality)
        if "mean_coherence" in quality:
            warnings.append(f"mean coherence {quality['mean_coherence']:.3f}")
        if "phase_complexity_ratio" in quality:
            warnings.append(f"phase complexity {quality['phase_complexity_ratio']:.3f}")
        if "damping_ratio" in quality:
            warnings.append(f"damping {quality['damping_ratio']:.4g}")
        rows.append({"record": entry.number, "frequency_hz": entry.frequency_hz, "source": entry.source_text,
                     "classification": STATUS_TEXT.get(entry.classification, entry.classification),
                     "modal_set": modal_set, "points": entry.point_count, "warnings": "; ".join(warnings)})
    return rows


def _proximity_rows(diagnostic) -> list[dict]:
    if diagnostic is None:
        return []
    rows = []
    for row in diagnostic.rows:
        rows.append({"experimental_record": row.experimental_mode, "experimental_hz": row.experimental_hz,
                     "nearest_fe_mode": row.nearest_fe_mode, "nearest_fe_hz": row.nearest_fe_hz,
                     "difference_to_nearest_percent": row.nearest_difference_percent,
                     "next_fe_mode": row.next_fe_mode, "next_fe_hz": row.next_fe_hz,
                     "difference_to_next_percent": row.next_difference_percent,
                     "range_note": diagnostic.range_note(row.experimental_hz), "mac": NOT_COMPUTED,
                     "pair": "NONE (not a mode pair)"})
    return rows


def _measured_dofs(experimental) -> str:
    if experimental is None:
        return "not available"
    axes = sorted({axis for entry in viewer.viewer_modes(experimental, viewer.EXPERIMENTAL) for axis in entry.axes})
    names = {"x": "U1", "y": "U2", "z": "U3"}
    return "/".join(names[axis] for axis in axes) or "none"


def build_diagnostic_report(loaded, *, abaqus_path: Optional[Path] = None, experimental_path: Optional[Path] = None,
                            images: Sequence[tuple[Path, str]] = (), cancel=None) -> DiagnosticReport:
    """The report of one analysis's loaded data (no comparison result); every value comes from the backend objects."""
    from services.guided_registration import REQUIRED_EVIDENCE

    abaqus, experimental, diagnostic = loaded.abaqus, loaded.experimental, loaded.frequency_diagnostic
    abaqus_path = Path(abaqus_path or getattr(loaded, "abaqus_path", None) or abaqus.source_path)
    experimental_path = experimental_path or getattr(loaded, "experimental_path", None)
    fe_file, exp_file = _identity_of(abaqus_path, cancel), _identity_of(experimental_path, cancel)
    try:
        geometry = dict(modal_dataset_geometry_identity(abaqus))
    except ValueError as error:
        geometry = {"sha256": None, "node_count": None, "note": str(error)}
    fe_rows, fe_range = _fe_rows(abaqus)
    experimental_rows = _experimental_rows(experimental, fe_range)
    metadata = dict(abaqus.metadata or {})
    is_manifest = abaqus_path.name.lower() == "manifest.json"
    fe_kind = "verified extraction manifest (the ODB itself was not read)" if is_manifest else "Abaqus ODB"
    status = "FREQUENCY_ONLY_DIAGNOSTIC" if experimental is not None else "EXPERIMENTAL_DATA_NOT_USABLE"
    banner = BANNER if experimental is not None else FE_ONLY_BANNER
    exp_source = (viewer.experimental_frequency_source(experimental) if experimental is not None
                  else f"not usable: {loaded.experimental_error}")
    sources = (
        ("FE source", f"{fe_kind}: {abaqus_path}"),
        ("FE file size / SHA-256", f"{fe_file['size']} bytes / {fe_file['sha256']}" + (
            f" ({fe_file['note']})" if fe_file["note"] else "")),
        ("ODB named by the extraction", str(metadata.get("odb_path") or metadata.get("odb_name") or "not recorded")),
        ("Abaqus step", str(metadata.get("step_name") or "not recorded")),
        ("FE geometry identity", f"{geometry.get('sha256')} ({geometry.get('node_count')} nodes)"),
        ("FE modes extracted", f"{len(fe_rows)}" + (f" ({fe_range[0]:.3f}–{fe_range[1]:.3f} Hz)" if fe_range
                                                    else "")),
        ("Experimental file", str(experimental_path) if experimental_path else "not recorded"),
        ("Experimental size / SHA-256", f"{exp_file['size']} bytes / {exp_file['sha256']}" + (
            f" ({exp_file['note']})" if exp_file["note"] else "")),
        ("Experimental source", exp_source),
        ("Modal set", experimental_rows[0]["modal_set"] if experimental_rows else "none"),
        ("Experimental records", str(len(experimental_rows))),
        ("Measured DOFs", _measured_dofs(experimental)),
        ("Physical registration", f"{loaded.registration} — not established"),
        ("MAC", NOT_COMPUTED),
    )
    reason = (diagnostic.reason if diagnostic is not None else loaded.reason) or "not recorded"
    passed = [f"Abaqus results loaded: {len(fe_rows)} FE modes" + (
        f", {fe_range[0]:.3f}–{fe_range[1]:.3f} Hz" if fe_range else ""),
        f"FE geometry identity computed: {geometry.get('sha256')} ({geometry.get('node_count')} nodes)"]
    not_passed = []
    if experimental is not None:
        passed.append(f"experimental file imported: {len(experimental_rows)} records ({exp_source})")
        not_passed.append(f"physical registration: {loaded.registration} — {reason}")
    else:
        not_passed.append(f"experimental import: {loaded.experimental_error or reason}")
        not_passed.append("physical registration: not attempted without usable experimental data")
    needed = tuple(text for _, text in REQUIRED_EVIDENCE) + (
        "or an accepted physical registration of exactly these files (FE geometry, experimental content, modal "
        "set); then 3. Mode shapes > Establish physical registration... gives the gated comparison with MAC",)
    out_of_range = [row for row in experimental_rows if "extracted FE mode" in row["warnings"]]
    provisional = sum(1 for row in experimental_rows if row["classification"].startswith("PROVISIONAL"))
    conclusions = (
        f"From Abaqus: {len(fe_rows)} FE modes with frequencies and FE mode shapes on the FE nodes.",
        f"From the experiment: {len(experimental_rows)} records" + (
            f", {provisional} of them PROVISIONAL FRF peaks (not validated fitted modes)" if provisional else "")
        + (f"; {len(out_of_range)} lie outside the extracted FE frequency range" if out_of_range else "") + ".",
        "Comparable only diagnostically: the two frequency lists and their nearest-frequency proximity.",
        "Not established: the physical registration, any mode pair, any MAC and any shape correspondence.",
        "Needed for a physical comparison: the independent physical evidence listed under 'What is needed'.",
        "No material constant is identified or implied by this report.",
    )
    provenance = {
        "schema": SCHEMA, "status": status, "registration": loaded.registration,
        "fe_file": fe_file, "experimental_file": exp_file, "fe_geometry_identity": geometry,
        "fe_kind": fe_kind, "extraction": {k: metadata.get(k) for k in ("odb_path", "odb_name", "step_name",
                                                                          "start_mode", "end_mode",
                                                                          "extraction_cache_reused")},
        "experimental_metadata": {k: (experimental.metadata.get(k) if experimental is not None else None)
                                  for k in ("modal_set_key", "modal_set_name", "mode_source", "frf_channel_count")},
        "frequency_only_diagnostic": None if diagnostic is None else diagnostic.status,
        "mac": NOT_COMPUTED, "mode_pairs": "NONE", "orientation": "NOT CHOSEN",
        "generated_by": "services.diagnostic_report from the loaded analysis data; no Abaqus, no new computation",
    }
    return DiagnosticReport(
        created=datetime.now().strftime("%Y-%m-%d %H:%M:%S"), status=status, banner=banner, sources=sources,
        reason=reason, checks_passed=tuple(passed), checks_not_passed=tuple(not_passed), needed=needed,
        fe_rows=tuple(fe_rows), fe_range_hz=fe_range, experimental_rows=tuple(experimental_rows),
        proximity_rows=tuple(_proximity_rows(diagnostic)),
        proximity_source="services.new_odb_comparison.frequency_only_diagnostic" if diagnostic is not None else "",
        conclusions=conclusions, provenance=provenance,
        images=tuple((str(path), caption) for path, caption in images if Path(path).exists()))


# ----------------------------------------------------------------------------- writing

def _atomic_write(destination: Path, write) -> Path:
    """Write to a temporary file next to the destination, then replace it (no partial report on failure)."""
    destination = Path(destination)
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        handle, temporary = tempfile.mkstemp(prefix=".~report_", suffix=destination.suffix, dir=destination.parent)
        os.close(handle)
    except OSError as error:
        raise ReportWriteError(f"the folder {destination.parent} is not writable: {error}") from error
    try:
        write(Path(temporary))
        os.replace(temporary, destination)
    except PermissionError as error:
        raise ReportWriteError(f"{destination.name} cannot be replaced; it may be open in another program "
                               f"(close it and export again): {error}") from error
    except OSError as error:
        raise ReportWriteError(f"{destination} cannot be written: {error}") from error
    finally:
        if os.path.exists(temporary):
            try:
                os.remove(temporary)
            except OSError:
                pass
    return destination


def _hz(value) -> str:
    return "—" if value is None else f"{value:.4f}"


def _percent(value) -> str:
    return "—" if value is None else f"{value:+.2f}"


def _short_note(note: str) -> str:
    return "below FE range" if note.startswith("below") else "above FE range" if note.startswith("above") else note


def _sections(report: DiagnosticReport) -> list[tuple[str, list[str]]]:
    width = 118
    wrap = lambda text, indent="": textwrap.wrap(text, width, initial_indent=indent,  # noqa: E731
                                                 subsequent_indent=" " * len(indent)) or [indent]
    sources = [line for item, value in report.sources for line in wrap(f"{item}: {value}", "")]
    reason = wrap(report.reason, "") + [""] + [line for text in report.checks_passed for line in wrap(text, "PASSED  ")]
    reason += [line for text in report.checks_not_passed for line in wrap(text, "NOT PASSED  ")]
    needed = [line for text in report.needed for line in wrap(text, "- ")]
    fe = [f"{'FE mode':>8}  {'Frequency, Hz':>14}  {'Points':>9}  {'Components':<11} Shape information"]
    fe += [f"{row['mode']:>8}  {row['frequency_hz']:>14.4f}  {row['points']:>9}  {row['components']:<11} "
           f"{row['shape']}"[:width] for row in report.fe_rows]
    if report.fe_range_hz:
        fe.append(f"Extracted FE frequency range: {report.fe_range_hz[0]:.4f} – {report.fe_range_hz[1]:.4f} Hz")
    exp = [f"{'Record':>6}  {'Frequency, Hz':>14}  Classification / warnings"]
    for row in report.experimental_rows:
        exp.append(f"{row['record']:>6}  {row['frequency_hz']:>14.4f}  {row['classification']}")
        if row["warnings"]:
            exp += wrap(row["warnings"], " " * 24)
    if report.experimental_rows:
        exp.append(f"Modal set: {report.experimental_rows[0]['modal_set']}")
    proximity = [NOT_PAIRING, "Differences are (f_FE − f_exp) / f_exp of the frequencies only: not a pair "
                              "error, MAC NOT_COMPUTED, no mode pair."]
    if report.fe_range_hz:
        proximity.append(f"'below / above FE range': the experimental frequency lies outside the extracted FE modes "
                         f"({report.fe_range_hz[0]:.2f}–{report.fe_range_hz[1]:.2f} Hz); extend the Abaqus mode "
                         "range for higher records.")
    proximity.append("")
    proximity.append(f"{'Exp':>4} {'Exp Hz':>11}  {'Nearest FE':>10} {'FE Hz':>11} {'Diff %':>8}  "
                     f"{'Next FE':>7} {'FE Hz':>11} {'Diff %':>8}  Note")
    for row in report.proximity_rows:
        proximity.append(f"{row['experimental_record']:>4} {_hz(row['experimental_hz']):>11}  "
                         f"{str(row['nearest_fe_mode'] or chr(8212)):>10} {_hz(row['nearest_fe_hz']):>11} "
                         f"{_percent(row['difference_to_nearest_percent']):>8}  "
                         f"{str(row['next_fe_mode'] or chr(8212)):>7} {_hz(row['next_fe_hz']):>11} "
                         f"{_percent(row['difference_to_next_percent']):>8}  {_short_note(row['range_note'])}")
    if not report.proximity_rows:
        proximity.append("No frequency proximity rows (no usable experimental frequencies).")
    conclusions = [line for text in report.conclusions for line in wrap(text, "- ")]
    limits = [line for text in LIMITS for line in wrap(text, "- ")]
    return [("Sources", sources), ("Why the analysis is diagnostic", reason),
            ("What is needed for a physical comparison with MAC", needed), ("FE modes (Abaqus)", fe),
            ("Experimental data", exp), ("Frequency proximity (diagnostic only)", proximity),
            ("Conclusion", conclusions), ("Limits", limits)]


def write_diagnostic_pdf(report: DiagnosticReport, destination: Path) -> Path:
    def write(path: Path) -> None:
        with PdfPages(path) as pdf:
            lines_per_page, page = 58, [0]

            def new_page(first=False):
                figure = Figure(figsize=(8.27, 11.69))
                FigureCanvasAgg(figure)
                page[0] += 1
                figure.text(0.06, 0.975, TITLE, fontsize=9, color="#555555")
                figure.text(0.94, 0.975, f"page {page[0]}", fontsize=9, color="#555555", ha="right")
                if first:
                    figure.text(0.06, 0.935, TITLE, fontsize=17, weight="bold")
                    figure.text(0.06, 0.905, report.banner, fontsize=11, weight="bold", color="#B00020")
                    figure.text(0.06, 0.885, f"Generated {report.created} · status {report.status} · MAC "
                                             f"{NOT_COMPUTED} · mode pairs NONE", fontsize=8)
                    return figure, 0.855
                figure.text(0.06, 0.955, report.banner, fontsize=8, weight="bold", color="#B00020")
                return figure, 0.935

            figure, y = new_page(first=True)
            step = 0.0148
            for heading, lines in _sections(report):
                if y < 0.08 + step * 4:
                    pdf.savefig(figure)
                    figure, y = new_page()
                figure.text(0.06, y, heading, fontsize=11, weight="bold")
                y -= step * 1.4
                for line in lines:
                    if y < 0.05:
                        pdf.savefig(figure)
                        figure, y = new_page()
                    figure.text(0.06, y, line, fontsize=6.6, family="monospace")
                    y -= step
                y -= step * 0.8
            pdf.savefig(figure)
            for image, caption in report.images:
                figure = Figure(figsize=(8.27, 11.69))
                FigureCanvasAgg(figure)
                axis = figure.add_axes((0.05, 0.12, 0.9, 0.78))
                axis.imshow(mpimg.imread(image))
                axis.axis("off")
                figure.text(0.06, 0.95, report.banner, fontsize=8, weight="bold", color="#B00020")
                figure.text(0.06, 0.06, caption, fontsize=8, wrap=True)
                pdf.savefig(figure)
            info = pdf.infodict()
            info["Title"] = TITLE
            info["Subject"] = report.banner

    return _atomic_write(destination, write)


HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="44546A")
WARN = Font(bold=True, color="B00020")


def _table(sheet, headers: Sequence[str], rows: Sequence[Sequence[Any]], start_row: int = 1, widths=()) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=column, value=title)
        cell.font, cell.fill = HEADER, HEADER_FILL
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=start_row + offset, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 40)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = width
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)


def write_diagnostic_xlsx(report: DiagnosticReport, destination: Path) -> Path:
    def write(path: Path) -> None:
        book = Workbook()
        summary = book.active
        summary.title = "Summary"
        summary["A1"], summary["A2"] = TITLE, report.banner
        summary["A1"].font, summary["A2"].font = Font(bold=True, size=14), WARN
        rows = [("Status", report.status), ("Generated", report.created), ("MAC", NOT_COMPUTED),
                ("Mode pairs", "NONE (physical mode pairing not established)"), *report.sources,
                ("Why diagnostic", report.reason), *(("Check passed", text) for text in report.checks_passed),
                *(("Check not passed", text) for text in report.checks_not_passed),
                *(("Needed for MAC", text) for text in report.needed),
                *(("Conclusion", text) for text in report.conclusions), *(("Limit", text) for text in LIMITS)]
        _table(summary, ("Item", "Value"), rows, start_row=4, widths=(30, 140))
        fe = book.create_sheet("FE Modes")
        _table(fe, ("FE mode", "Frequency, Hz", "Points", "Components with data", "Shape information"),
               [(r["mode"], r["frequency_hz"], r["points"], r["components"], r["shape"]) for r in report.fe_rows],
               widths=(10, 16, 12, 22, 60))
        experimental = book.create_sheet("Experimental Data")
        _table(experimental, ("Record", "Frequency, Hz", "Classification", "Source", "Modal set", "Points",
                              "Diagnostic warnings"),
               [(r["record"], r["frequency_hz"], r["classification"], r["source"], r["modal_set"], r["points"],
                 r["warnings"]) for r in report.experimental_rows], widths=(9, 16, 44, 60, 22, 9, 70))
        proximity = book.create_sheet("Frequency Diagnostic")
        proximity["A1"] = NOT_PAIRING + " Differences are (f_FE - f_exp) / f_exp of frequencies only; MAC NOT_COMPUTED."
        proximity["A1"].font = WARN
        _table(proximity, ("Experimental record", "Experimental, Hz", "Nearest FE mode", "Nearest FE, Hz",
                           "Difference to nearest, %", "Next FE mode", "Next FE, Hz", "Difference to next, %",
                           "Range note", "MAC", "Mode pair"),
               [(r["experimental_record"], r["experimental_hz"], r["nearest_fe_mode"], r["nearest_fe_hz"],
                 r["difference_to_nearest_percent"], r["next_fe_mode"], r["next_fe_hz"],
                 r["difference_to_next_percent"], r["range_note"], r["mac"], r["pair"])
                for r in report.proximity_rows], start_row=3, widths=(12, 16, 14, 16, 18, 12, 16, 18, 50, 14, 22))
        provenance = book.create_sheet("Provenance")

        def flat(prefix, value):
            if isinstance(value, Mapping):
                for key, item in value.items():
                    yield from flat(f"{prefix}.{key}" if prefix else str(key), item)
            else:
                yield prefix, "" if value is None else str(value)

        _table(provenance, ("Field", "Value"), list(flat("", report.provenance)), widths=(48, 120))
        book.save(path)

    return _atomic_write(destination, write)
