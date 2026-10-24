"""APP-E2E.3: preparation for repeat measurements (synthetic data, no process, no Abaqus).

- P-2: FE rigid-body modes are separate diagnostic records by the comparator's own shape criterion (no frequency
  threshold); the first elastic FE mode and every experimental mode stay available;
- P-3: the 11. Scaling tab wraps to its real width and nothing is cut, also at 150 % (Tk scaling 2.0);
- SETUP_UNDER_REVIEW: a person's note bound to the series identity (content and modal set), mandatory reason,
  source and experiment id, append-only, shown in the GUI and the diagnostic export, and it changes no data,
  frequency, shape, MAC, pairing or result.
"""

from __future__ import annotations

import copy
import gc
import hashlib
import json
from pathlib import Path
import shutil
import sys
import tempfile
import time
import unittest
from unittest import mock

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

import synthetic_plate_support as S  # noqa: E402
from modal_core import ModalDataset, ModeShape  # noqa: E402
from services import diagnostic_report as dr  # noqa: E402
from services import experimental_series_review as review  # noqa: E402
from services import modal_data_viewer as viewer  # noqa: E402
from services import new_odb_comparison as new  # noqa: E402
from test_modal_data_viewer import fe_plate, load_unv  # noqa: E402


def _fingerprint(path: Path) -> tuple:
    stat = path.stat()
    return hashlib.sha256(path.read_bytes()).hexdigest(), stat.st_size, stat.st_mtime_ns


class _Files(unittest.TestCase):
    def setUp(self) -> None:
        self.folder = Path(tempfile.mkdtemp(prefix="e2e3_"))
        self.addCleanup(shutil.rmtree, self.folder, True)
        self.fe = fe_plate()
        self.unv = S.write_experiment(self.folder / "plate.unv", S.E0, S.G0)
        self.store = self.folder / "review" / "series_review.json"
        patcher = mock.patch("subprocess.Popen", side_effect=AssertionError("no process may start"))
        patcher.start()
        self.addCleanup(patcher.stop)

    def loaded(self, experimental_path=None, fe=None):
        experimental = load_unv(experimental_path or self.unv)
        fe = fe or self.fe
        loaded = viewer.LoadedModalData(fe, experimental, abaqus_path=self.folder / "plate.odb",
                                        experimental_path=experimental_path or self.unv,
                                        registration=viewer.AMBIGUOUS, reason="8 orientations")
        (self.folder / "plate.odb").write_bytes(b"synthetic ODB stand-in; never opened")
        loaded.frequency_diagnostic = new.frequency_only_diagnostic(fe, experimental, "8 orientations")
        return loaded


# ----------------------------------------------------------------------------- P-2

class RigidBodyRecordTests(_Files):
    def test_rigid_body_modes_are_separate_diagnostic_records_and_the_first_elastic_mode_stays(self):
        experimental = load_unv(self.unv)
        diagnostic = new.frequency_only_diagnostic(self.fe, experimental, "x")
        self.assertEqual([n for n, *_ in diagnostic.rigid_fe_modes], [1, 2, 3])
        self.assertEqual([n for n, _ in diagnostic.fe_modes], [4, 5, 6, 7, 8, 9])  # first elastic FE mode kept
        self.assertTrue(all(row.nearest_fe_mode >= 4 and row.next_fe_mode >= 4 for row in diagnostic.rows))
        data = diagnostic.to_dict()
        self.assertEqual([r["mode"] for r in data["rigid_body_fe_modes"]], [1, 2, 3])
        self.assertIn("no frequency threshold", data["rigid_body_criterion"])
        self.assertIn("FE RIGID-BODY MODES", "\n".join(diagnostic.text_lines()))
        labels = [mode.label() for mode in viewer.viewer_modes(self.fe, viewer.FE)]
        self.assertEqual(len(labels), 9)  # nothing removed from the FE list
        self.assertTrue(all("rigid-body" in label for label in labels[:3]))
        self.assertFalse(any("rigid-body" in label for label in labels[3:]))

    def test_the_criterion_is_the_shape_not_the_frequency(self):
        ids = np.array([f"P-1.{i}" for i in range(1, 5)], dtype=object)
        xyz = np.array([[0.0, 0.0, 0.0], [1.0, 0.0, 0.0], [0.0, 1.0, 0.0], [1.0, 1.0, 0.0]])
        bending = np.zeros((4, 3))
        bending[:, 2] = [1.0, -1.0, -1.0, 1.0]  # a twist: not a rigid-body motion
        translation = np.zeros((4, 3))
        translation[:, 2] = 1.0
        fe = ModalDataset("fe", Path("x.odb"), [ModeShape(1, 0.0004, ids, xyz, bending),
                                                 ModeShape(2, 55.0, ids, xyz, translation)])
        residuals = viewer.classify_rigid_body_fe_modes(fe)
        self.assertGreater(residuals[1], viewer.rigid_body_threshold())  # near zero Hz, yet elastic
        self.assertLessEqual(residuals[2], viewer.rigid_body_threshold())  # 55 Hz, yet rigid-body
        experimental = ModalDataset("exp", Path("e.unv"), [ModeShape(1, 0.0005, ids, xyz, bending)])
        diagnostic = new.frequency_only_diagnostic(fe, experimental, "x")
        self.assertEqual(diagnostic.fe_modes, ((1, 0.0004),))
        self.assertEqual(diagnostic.experimental_modes, ((1, 0.0005),))  # experimental mode 1 is never removed

    def test_a_recorded_classification_is_reused_never_recomputed_on_the_gui_thread(self):
        viewer.classify_rigid_body_fe_modes(self.fe)
        with mock.patch("quality_control._rigid_body_residual_fraction", side_effect=AssertionError("recomputed")):
            viewer.viewer_modes(self.fe, viewer.FE)
            viewer.classify_rigid_body_fe_modes(self.fe)
        fresh = fe_plate()  # never classified: the viewer does not compute, it only shows plain FE modes
        with mock.patch("quality_control._rigid_body_residual_fraction", side_effect=AssertionError("computed")):
            entries = viewer.viewer_modes(fresh, viewer.FE)
        self.assertTrue(all(entry.classification == viewer.FE_MODE for entry in entries))

    def test_the_diagnostic_report_lists_rigid_body_modes_separately(self):
        report = dr.build_diagnostic_report(self.loaded(), series_review_store=self.store)
        rows = {row["mode"]: row for row in report.fe_rows}
        self.assertEqual(len(rows), 9)
        self.assertTrue(rows[1]["rigid_body"] and rows[1]["classification"].startswith("RIGID-BODY"))
        self.assertEqual(rows[4]["classification"], "elastic FE mode")
        self.assertEqual(report.fe_range_hz, (60.0, 230.0))  # the elastic range
        self.assertEqual([row["record"] for row in report.experimental_rows][:1], [1])
        workbook = dr.write_diagnostic_xlsx(report, self.folder / "r.xlsx")
        from openpyxl import load_workbook

        sheet = load_workbook(workbook, read_only=True)["FE Modes"]
        values = list(sheet.iter_rows(values_only=True))
        self.assertEqual(values[0][-1], "Classification")
        self.assertTrue(str(values[1][-1]).startswith("RIGID-BODY"))


# ----------------------------------------------------------------------------- SETUP_UNDER_REVIEW

class SeriesReviewTests(_Files):
    def identity(self, path=None, modal_set="polymax"):
        return review.series_identity(path or self.unv, modal_set)

    def mark(self, identity, **overrides):
        values = dict(reason="suspension may have influenced the first mode", source="owner, 2026-10-25",
                      experiment_id="PLATE session 1")
        values.update(overrides)
        return review.mark_setup_under_review(self.store, identity, **values)

    def test_reason_source_and_experiment_id_are_required(self):
        identity = self.identity()
        for missing in ("reason", "source", "experiment_id"):
            with self.subTest(missing=missing), self.assertRaises(review.SeriesReviewError):
                self.mark(identity, **{missing: "  "})
        self.assertIsNone(review.active_mark(self.store, identity))

    def test_the_mark_belongs_to_the_series_content_and_modal_set_not_to_the_file_name(self):
        identity = self.identity()
        self.mark(identity)
        renamed = self.folder / "renamed_copy.unv"
        shutil.copyfile(self.unv, renamed)
        self.assertIsNotNone(review.active_mark(self.store, self.identity(renamed)))  # same content
        repeat = S.write_experiment(self.folder / "plate.unv.repeat", S.E0 * 1.01, S.G0)  # a new series
        self.assertIsNone(review.active_mark(self.store, self.identity(repeat)))
        self.assertIsNone(review.active_mark(self.store, self.identity(modal_set="another-set")))
        self.assertEqual(review.mark_text(None), review.NO_MARK)
        self.assertIn("not a statement that the series", review.NO_MARK)

    def test_withdrawal_is_recorded_and_the_log_is_append_only(self):
        identity = self.identity()
        first = self.mark(identity)
        with self.assertRaises(review.SeriesReviewError):
            review.withdraw_mark(self.store, identity, reason="", source="owner")
        review.withdraw_mark(self.store, identity, reason="repeated with the original suspension", source="owner")
        self.assertIsNone(review.active_mark(self.store, identity))
        entries = json.loads(self.store.read_text(encoding="utf-8"))["entries"]
        self.assertEqual([entry["action"] for entry in entries], [review.SET, review.WITHDRAWN])
        self.assertEqual(entries[0], first)  # never rewritten
        with self.assertRaises(review.SeriesReviewError):
            review.withdraw_mark(self.store, identity, reason="again", source="owner")

    def test_the_experimental_file_is_only_read(self):
        before = _fingerprint(self.unv)
        identity = self.identity()
        self.mark(identity)
        review.withdraw_mark(self.store, identity, reason="done", source="owner")
        self.assertEqual(_fingerprint(self.unv), before)

    def test_a_mark_changes_no_frequency_shape_mac_or_pairing(self):
        experimental = load_unv(self.unv)
        shapes_before = [copy.deepcopy(np.asarray(mode.vectors)) for mode in experimental.sorted_modes()]
        before = new.frequency_only_diagnostic(fe_plate(), experimental, "x").to_dict()
        self.mark(self.identity())
        after = new.frequency_only_diagnostic(fe_plate(), load_unv(self.unv), "x").to_dict()
        self.assertEqual(before, after)
        self.assertEqual(after["mac"], "NOT_COMPUTED")
        self.assertEqual(after["mode_pairs"], "NONE (physical modal correspondence not verified)")
        self.assertEqual([r["experimental_mode"] for r in after["nearest_frequency_rows"]][:1], [1])
        for old, mode in zip(shapes_before, load_unv(self.unv).sorted_modes()):
            np.testing.assert_array_equal(old, np.asarray(mode.vectors))

    def test_the_diagnostic_export_states_the_mark_or_its_absence(self):
        report = dr.build_diagnostic_report(self.loaded(), series_review_store=self.store)
        sources = dict(report.sources)
        self.assertEqual(sources["Experimental series review"], review.NO_MARK)
        self.assertIsNone(report.provenance["experimental_series_review"])
        self.mark(self.identity())
        report = dr.build_diagnostic_report(self.loaded(), series_review_store=self.store)
        self.assertTrue(dict(report.sources)["Experimental series review"].startswith(review.SETUP_UNDER_REVIEW))
        self.assertEqual(report.provenance["experimental_series_review"]["experiment_id"], "PLATE session 1")
        self.assertTrue(any("note on the test conditions only" in text for text in report.conclusions))
        self.assertEqual(report.status, "FREQUENCY_ONLY_DIAGNOSTIC")  # unchanged by the note
        pdf = dr.write_diagnostic_pdf(report, self.folder / "r.pdf")
        import pypdf

        text = "\n".join(page.extract_text() for page in pypdf.PdfReader(str(pdf)).pages)
        self.assertIn(review.SETUP_UNDER_REVIEW, text)


# ----------------------------------------------------------------------------- GUI

class _Window(_Files):
    def application(self, tk_scaling=None, size=None):
        try:
            import tkinter as tk

            root = tk.Tk()
        except Exception as error:  # noqa: BLE001
            self.skipTest(f"Tk is unavailable: {error}")
        from test_app_p1_stale_results import _App

        self.addCleanup(gc.collect)  # the window's objects are freed on the main thread
        self.addCleanup(lambda: root.destroy() if root.winfo_exists() else None)
        self.addCleanup(_App.cancel_after_jobs, root)
        self.addCleanup(_App.join_background_threads)  # series identification / modal-set discovery threads
        if tk_scaling:
            before = float(root.tk.call("tk", "scaling"))
            # Tk keeps the scaling per display, shared by every window of the process: restore it for later tests
            self.addCleanup(lambda: root.tk.call("tk", "scaling", before) if root.winfo_exists() else None)
            root.tk.call("tk", "scaling", tk_scaling)
        import main
        import ui_workflow

        disabled = {"autosave_enabled": False, "restore_on_startup": False, "show_recovery_notice": False}
        with mock.patch.object(ui_workflow, "load_recovery_preferences", return_value=disabled), \
                mock.patch.object(ui_workflow, "discover_abaqus_installations", return_value=()):
            application = main.app.ModalComparatorApp(root)
        application._series_review_store = self.store
        application.mode_viewer.threaded = False
        if size:
            root.geometry(f"{size[0]}x{size[1]}+0+0")
        self.root = root
        self.pump(0.3)
        return application

    def pump(self, seconds):
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            self.root.update()
            time.sleep(0.02)

    def wait(self, condition, seconds=30):
        end = time.monotonic() + seconds
        while time.monotonic() < end and not condition():
            self.root.update()
            time.sleep(0.02)
        self.assertTrue(condition())


class SeriesReviewWindowTests(_Window):
    def test_the_files_tab_shows_the_note_of_exactly_the_selected_series(self):
        application = self.application()
        application.experimental_path.set(str(self.unv))
        self.wait(lambda: application._series_identity is not None
                  and application._modal_set_discovery_state != "loading")
        self.wait(lambda: application._series_identity is not None)
        self.assertEqual(application.series_review_text.get(), review.NO_MARK)
        import series_review_ui

        window = series_review_ui.open_series_review_dialog(application)
        window.fields["experiment_id"].set("PLATE session 1")
        window.fields["reason"].set("suspension changed")
        window.fields["source"].set("owner")
        window.mark_button.invoke()
        self.assertTrue(application.series_review_text.get().startswith(review.SETUP_UNDER_REVIEW))
        window.destroy()
        application.tabs.select(application.input_tab)
        self.pump(0.8)
        label, button = application.series_review_label, application.series_review_button
        self.assertLessEqual(label.winfo_reqheight(), label.winfo_height() + 1)  # every line shown
        self.assertLessEqual(label.winfo_rootx() + label.winfo_width(), button.winfo_rootx())  # never under the button
        self.assertLessEqual(int(label.cget("wraplength")), label.winfo_width())
        other = S.write_experiment(self.folder / "repeat.unv", S.E0 * 1.02, S.G0)
        application.experimental_path.set(str(other))  # a repeat measurement: another series, no note
        self.wait(lambda: application._series_identity is not None
                  and application._series_identity.sha256 != hashlib.sha256(self.unv.read_bytes()).hexdigest())
        self.assertEqual(application.series_review_text.get(), review.NO_MARK)
        self.assertIsNone(getattr(application, "_analysis_inputs", None))  # never an analysis input


class ScalingTabTests(_Window):
    def assert_tab_fits(self, application):
        from tkinter import ttk

        application.tabs.select(application.scaling_tab)
        self.pump(1.0)
        panes = application.scaling_panes
        bottom, right = panes.winfo_rooty() + panes.winfo_height(), panes.winfo_rootx() + panes.winfo_width()
        stack = [panes]
        while stack:
            widget = stack.pop()
            stack.extend(widget.winfo_children())
            if not widget.winfo_ismapped() or widget is panes:
                continue
            self.assertLessEqual(widget.winfo_rooty() + widget.winfo_height(), bottom + 1, str(widget))
            self.assertLessEqual(widget.winfo_rootx() + widget.winfo_width(), right + 1, str(widget))
            if isinstance(widget, ttk.Label) and str(widget.cget("text")):
                self.assertLessEqual(widget.winfo_reqheight(), widget.winfo_height() + 1, str(widget))
            if isinstance(widget, (ttk.Button, ttk.Spinbox)):
                self.assertGreaterEqual(widget.winfo_width(), widget.winfo_reqwidth() - 2, str(widget))
        caveat = application.scaling_caveat
        self.assertGreater(int(caveat.cget("wraplength")), 0.6 * caveat.master.winfo_width())  # not a narrow column

    def test_scaling_tab_at_100_percent_and_minimum_window(self):
        from ui_policy import MINIMUM_WINDOW_SIZE

        self.assert_tab_fits(self.application(size=MINIMUM_WINDOW_SIZE))

    def test_scaling_tab_at_150_percent(self):
        from ui_policy import MINIMUM_WINDOW_SIZE

        for size in ((1680, 1000), MINIMUM_WINDOW_SIZE):
            with self.subTest(size=size):
                application = self.application(tk_scaling=2.0, size=size)
                application.scaling_summary.configure(text="AVAILABLE: " + "a descriptive summary line " * 12)
                self.assert_tab_fits(application)
                canvas = application.scaling_canvas
                if application.scaling_panes.winfo_height() > canvas.winfo_height() + 1:
                    self.assertTrue(application.scaling_vscroll.winfo_ismapped())  # reachable by scrolling

    def test_a_hidden_tab_is_not_wrapped_to_a_narrow_column(self):
        application = self.application(size=(1680, 1000))
        self.pump(0.5)  # the responsive layout runs while the Scaling tab is still hidden
        self.assertNotEqual(int(application.scaling_caveat.cget("wraplength")), 280)
        self.assert_tab_fits(application)


if __name__ == "__main__":
    unittest.main()
