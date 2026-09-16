from pathlib import Path
import sys
import tempfile
import unittest

import numpy as np
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from modal_core import ModalDataset, ModeShape, compare_modal_datasets, modal_assurance_criterion
from reporting import export_excel, export_pdf, render_pair_images


class ModalCoreTests(unittest.TestCase):
    def synthetic_result(self):
        x, y = np.meshgrid(np.linspace(0.0, 500.0, 13), np.linspace(0.0, 510.0, 11))
        abaqus_coordinates = np.column_stack((x.ravel(), y.ravel(), np.zeros(x.size)))
        # A complete, evenly-spaced rectangular grid is invariant under a
        # 180-degree rotation and both axis mirrors regardless of aspect
        # ratio; nudge one point off that exact symmetry so the point set
        # has exactly one geometrically admissible registration, for a
        # reason unrelated to what these tests verify.
        abaqus_coordinates[0] += np.array([0.37, -0.21, 0.0])
        abaqus_ids = np.arange(len(abaqus_coordinates))

        experimental_indices = np.arange(0, len(abaqus_coordinates), 4)
        selected = abaqus_coordinates[experimental_indices]
        true_rotation = np.array(
            [
                [0.0, -1.0, 0.0],
                [1.0, 0.0, 0.0],
                [0.0, 0.0, 1.0],
            ]
        )
        center = (abaqus_coordinates.min(axis=0) + abaqus_coordinates.max(axis=0)) / 2.0
        experimental_coordinates = (selected - center) @ true_rotation * 0.001 + np.array([2.0, 3.0, 0.0])

        abaqus_modes = []
        experimental_modes = []
        # Modes 2 and 3 are deliberately close and reverse their frequency order,
        # but every physical pair remains inside the reviewed admissibility gates.
        frequencies = [105.0, 148.0, 154.0]
        experimental_order = [1, 2, 3]
        experimental_frequencies = [106.5, 155.0, 150.0]

        for index, frequency in enumerate(frequencies, start=1):
            vectors = np.zeros_like(abaqus_coordinates)
            vectors[:, 2] = (
                np.sin(index * np.pi * abaqus_coordinates[:, 0] / 500.0)
                * np.sin((index + 1) * np.pi * abaqus_coordinates[:, 1] / 510.0)
            )
            abaqus_modes.append(
                ModeShape(
                    number=index + 5,
                    frequency_hz=frequency,
                    node_ids=abaqus_ids,
                    coordinates=abaqus_coordinates,
                    vectors=vectors,
                )
            )
            experimental_modes.append(
                ModeShape(
                    number=experimental_order[index - 1],
                    frequency_hz=experimental_frequencies[index - 1],
                    node_ids=np.arange(len(experimental_indices)),
                    coordinates=experimental_coordinates,
                    vectors=vectors[experimental_indices] @ true_rotation,
                )
            )

        return compare_modal_datasets(
            ModalDataset("Abaqus", Path("model.odb"), abaqus_modes),
            ModalDataset("Testlab", Path("panel.unv"), experimental_modes),
        )

    def test_mac_identical_vectors(self):
        vector = np.array([1.0, -2.0, 3.0, 0.5])
        self.assertAlmostEqual(modal_assurance_criterion(vector, -4.0 * vector), 1.0)

    def test_complex_mac_is_invariant_to_global_phase(self):
        phi = np.array([1.0 + 2.0j, -3.0 + 0.5j, 2.0 - 4.0j])
        theta = 0.73
        psi = np.exp(1.0j * theta) * phi

        self.assertAlmostEqual(modal_assurance_criterion(phi, psi), 1.0)

    def test_complex_mac_uses_hermitian_inner_product(self):
        phi = np.array([1.0 + 1.0j, 2.0 - 1.0j, -0.5 + 2.0j])
        psi = np.array([2.0 - 1.0j, -1.0 + 0.5j, 1.0 + 3.0j])
        denominator = np.vdot(phi, phi).real * np.vdot(psi, psi).real
        hermitian_expected = abs(np.vdot(phi, psi)) ** 2 / denominator
        plain_transpose_result = abs(np.dot(phi, psi)) ** 2 / denominator

        self.assertAlmostEqual(
            modal_assurance_criterion(phi, psi), hermitian_expected
        )
        self.assertGreater(
            abs(hermitian_expected - plain_transpose_result),
            0.1,
        )

    def test_complex_mac_is_invariant_to_nonzero_complex_scaling(self):
        phi = np.array([1.0 + 1.0j, 2.0 - 1.0j, -0.5 + 2.0j])
        psi = np.array([2.0 - 1.0j, -1.0 + 0.5j, 1.0 + 3.0j])
        baseline = modal_assurance_criterion(phi, psi)

        scaled = modal_assurance_criterion((2.0 - 3.0j) * phi, (-0.5 + 1.25j) * psi)

        self.assertAlmostEqual(scaled, baseline)

    def test_complex_mac_zero_norm_returns_none(self):
        zero = np.zeros(3, dtype=complex)
        nonzero = np.array([1.0 + 2.0j, -3.0j, 0.5 - 0.25j])

        self.assertIsNone(modal_assurance_criterion(zero, nonzero))
        self.assertIsNone(modal_assurance_criterion(nonzero, zero))

    def test_geometry_alignment_and_mode_matching(self):
        result = self.synthetic_result()
        self.assertGreaterEqual(result.geometry.matched_fraction, 0.95)
        self.assertLess(result.geometry.normalized_rms_distance, 1e-8)
        self.assertEqual(len(result.pairs), 3)
        available_mac = [pair.mac for pair in result.pairs if pair.mac is not None]
        self.assertTrue(available_mac)
        self.assertGreater(min(available_mac), 0.98)
        self.assertTrue(any(pair.order_changed for pair in result.pairs))

    def test_report_generation(self):
        result = self.synthetic_result()
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            render_pair_images(result.pairs[0], directory_path / "images")
            excel_path = export_excel(result, directory_path / "report.xlsx", directory_path / "report_images")
            pdf_path = export_pdf(result, directory_path / "report.pdf", directory_path / "report_images")
            self.assertGreater(excel_path.stat().st_size, 1000)
            self.assertGreater(pdf_path.stat().st_size, 1000)

    def test_export_excel_base_writes_the_documented_sheets_and_values(self):
        """Characterization test for the un-patched reporting.export_excel.

        Locks in the sheet list, headers, and summary values so a future
        refactor of the install_* patch chain (see ROADMAP.md Stage 4) can be
        checked against this instead of just "did it crash".
        """
        result = self.synthetic_result()
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            excel_path = export_excel(
                result, directory_path / "report.xlsx", directory_path / "report_images"
            )
            workbook = load_workbook(excel_path)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "Summary",
                        "Mode Comparison",
                        "MAC Matrix",
                        "Geometry Match",
                        "Abaqus History",
                    ],
                )

                summary = workbook["Summary"]
                self.assertEqual(summary["A1"].value, "Abaqus Modal Comparator")
                self.assertEqual(summary["A5"].value, "Matched mode pairs")
                self.assertEqual(summary["B5"].value, len(result.pairs))
                self.assertEqual(summary["A10"].value, "Mean frequency error")
                self.assertAlmostEqual(
                    summary["B10"].value,
                    float(np.mean([pair.frequency_error_percent for pair in result.pairs])) / 100.0,
                    places=8,
                )
                mac_values = [pair.mac for pair in result.pairs if pair.mac is not None]
                self.assertTrue(mac_values, "fixture must include at least one calculable MAC")
                self.assertEqual(summary["A11"].value, "Mean MAC")
                self.assertAlmostEqual(
                    summary["B11"].value, float(np.mean(mac_values)), places=8
                )

                comparison = workbook["Mode Comparison"]
                self.assertEqual(
                    [cell.value for cell in comparison[1]],
                    [
                        "Abaqus mode",
                        "Experimental mode",
                        "Abaqus frequency, Hz",
                        "Experimental frequency, Hz",
                        "Frequency error, %",
                        "MAC",
                        "Order changed",
                        "Mapped points",
                        "Status",
                    ],
                )
                self.assertEqual(comparison.max_row, len(result.pairs) + 1)

                mac_sheet = workbook["MAC Matrix"]
                self.assertEqual(
                    [cell.value for cell in mac_sheet[1]],
                    ["Abaqus / Experiment", *result.experimental_mode_numbers],
                )
            finally:
                workbook.close()

    def test_export_pdf_base_produces_one_page_per_summary_item_and_pair(self):
        """Characterization test for the un-patched reporting.export_pdf page count:
        1 title page + 2 diagnostic pages (frequency, MAC) + 1 page per matched pair.
        """
        result = self.synthetic_result()
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            pdf_path = export_pdf(
                result, directory_path / "report.pdf", directory_path / "report_images"
            )
            reader = PdfReader(str(pdf_path))
            self.assertEqual(len(reader.pages), 3 + len(result.pairs))


if __name__ == "__main__":
    unittest.main()
