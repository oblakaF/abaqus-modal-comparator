"""CORE-E2E-D: diagnostic PDF / Excel report of an analysis without a physical registration (no Abaqus).

The report is read back (pypdf text, openpyxl cells) and compared with the backend objects it was built from: FE and
experimental frequencies, classifications, the frequency-only diagnostic rows, range notes and file identities. It
never carries a MAC value, a mode pair or a pair error; a write failure leaves no partial file; the export buttons
keep the comparison report unchanged and export the diagnostic only when there is no comparison result.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

import numpy as np
from openpyxl import load_workbook
import pypdf

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

import frf_unv_support as F  # noqa: E402
import synthetic_plate_support as S  # noqa: E402
from modal_core import ModalDataset  # noqa: E402
from scientific_state import modal_dataset_geometry_identity  # noqa: E402
from services import diagnostic_report as dr  # noqa: E402
from services import modal_data_viewer as viewer  # noqa: E402
from services import new_odb_comparison as new  # noqa: E402
from services.archived_sensitivity_evidence import fresh_content_identity  # noqa: E402
from test_modal_data_viewer import load_unv  # noqa: E402

MAC_VALUE = re.compile(r"MAC[^A-Za-z\n]{0,12}0\.\d")


def pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() for page in pypdf.PdfReader(str(path)).pages)


def sheet_rows(path: Path, name: str) -> list[tuple]:
    book = load_workbook(path, read_only=True)
    try:
        return [row for row in book[name].iter_rows(values_only=True)]
    finally:
        book.close()


def sheet_names(path: Path) -> list[str]:
    book = load_workbook(path, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def fe_with_instances() -> ModalDataset:
    """The synthetic plate with INSTANCE:label node ids (an FE geometry identity exists)."""
    fake = S.FakeAbaqus()
    fake.node_id_format = "PLATE-1:{}"
    folder = Path(tempfile.mkdtemp(prefix="dr_fe_"))
    odb = folder / "plate.odb"
    odb.write_text(json.dumps({"E1": S.E0, "E2": S.E0, "G12": S.G0}))
    dataset = fake.extract(odb, folder, 9)
    shutil.rmtree(folder, ignore_errors=True)
    return dataset


class _Case(unittest.TestCase):
    def setUp(self) -> None:
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.folder = Path(folder.name)
        self.fe = fe_with_instances()
        self.odb = self.folder / "plate.odb"
        self.odb.write_bytes(b"synthetic ODB stand-in; never opened")
        patcher = mock.patch("subprocess.Popen", side_effect=AssertionError("no process may start"))
        patcher.start()
        self.addCleanup(patcher.stop)

    def loaded(self, experimental_path: Path | None, fe: ModalDataset | None = None, modal_set=None):
        fe = fe or self.fe
        if experimental_path is None:
            return viewer.LoadedModalData(fe, None, experimental_error="no modal or FRF record in the file",
                                          abaqus_path=self.odb, registration=viewer.NOT_ESTABLISHED,
                                          reason="the experimental file gave no usable modal or FRF data")
        experimental = load_unv(experimental_path, modal_set=modal_set)
        loaded = viewer.LoadedModalData(fe, experimental, abaqus_path=self.odb, experimental_path=experimental_path,
                                        registration=viewer.AMBIGUOUS, reason="8 orientations")
        loaded.frequency_diagnostic = new.frequency_only_diagnostic(
            fe, experimental, "8 physically distinct orientations are equally plausible from geometry alone, and no "
                              "accepted physical registration could be reused: synthetic test file")
        return loaded

    def written(self, loaded):
        report = dr.build_diagnostic_report(loaded)
        return (report, dr.write_diagnostic_pdf(report, self.folder / "report.pdf"),
                dr.write_diagnostic_xlsx(report, self.folder / "report.xlsx"))


class FittedDiagnosticTests(_Case):
    def setUp(self) -> None:
        super().setUp()
        self.unv = S.write_experiment(self.folder / "fitted.unv", S.E0, S.G0)
        self.loaded_data = self.loaded(self.unv)
        self.report, self.pdf, self.xlsx = self.written(self.loaded_data)

    def test_frequencies_and_classifications_come_from_the_loaded_data(self):
        fe_hz = [m.frequency_hz for m in self.fe.sorted_modes()]
        exp_hz = [m.frequency_hz for m in self.loaded_data.experimental.sorted_modes()]
        self.assertEqual([r["frequency_hz"] for r in self.report.fe_rows], fe_hz)
        self.assertEqual([r["frequency_hz"] for r in self.report.experimental_rows], exp_hz)
        self.assertTrue(all(r["classification"].startswith("FITTED MODE") for r in self.report.experimental_rows))
        rows = self.loaded_data.frequency_diagnostic.rows
        self.assertEqual([(r["experimental_record"], r["nearest_fe_mode"], r["difference_to_nearest_percent"])
                          for r in self.report.proximity_rows],
                         [(r.experimental_mode, r.nearest_fe_mode, r.nearest_difference_percent) for r in rows])
        self.assertEqual({r["mac"] for r in self.report.proximity_rows}, {dr.NOT_COMPUTED})

    def test_workbook_sheets_and_values(self):
        self.assertEqual(sheet_names(self.xlsx), ["Summary", "FE Modes", "Experimental Data", "Frequency Diagnostic",
                                           "Provenance"])
        fe = sheet_rows(self.xlsx, "FE Modes")[1:]
        self.assertEqual([row[1] for row in fe], [m.frequency_hz for m in self.fe.sorted_modes()])
        diagnostic = sheet_rows(self.xlsx, "Frequency Diagnostic")
        self.assertIn(dr.NOT_PAIRING, diagnostic[0][0])
        header = diagnostic[2]
        body = diagnostic[3:]
        self.assertEqual({row[header.index("MAC")] for row in body}, {dr.NOT_COMPUTED})
        self.assertEqual({row[header.index("Mode pair")] for row in body}, {"NONE (not a mode pair)"})
        summary = dict((row[0], row[1]) for row in sheet_rows(self.xlsx, "Summary")[3:] if row[0])
        self.assertEqual(summary["MAC"], dr.NOT_COMPUTED)
        self.assertEqual(summary["Status"], "FREQUENCY_ONLY_DIAGNOSTIC")
        self.assertIn("could be reused: synthetic test file", summary["Why diagnostic"])

    def test_pdf_states_the_diagnostic_and_carries_no_mac_or_pair(self):
        text = pdf_text(self.pdf)
        for required in (dr.TITLE, dr.BANNER, dr.NOT_PAIRING, "MAC NOT_COMPUTED", "mode pairs NONE",
                         "What is needed for a physical comparison with MAC", "NOT PASSED  physical registration"):
            self.assertIn(required, text)
        self.assertIsNone(MAC_VALUE.search(text))
        for forbidden in ("Excellent match", "Good match", "Matched pairs", "Overlay"):
            self.assertNotIn(forbidden, text)
        workbook_text = " ".join(str(cell) for name in sheet_names(self.xlsx)
                                 for row in sheet_rows(self.xlsx, name) for cell in row if cell is not None)
        self.assertIsNone(MAC_VALUE.search(workbook_text))
        self.assertNotIn("E1 =", text + workbook_text)

    def test_provenance_identities_are_the_files(self):
        provenance = self.report.provenance
        self.assertEqual(provenance["experimental_file"]["sha256"], fresh_content_identity(self.unv)["sha256"])
        self.assertEqual(provenance["fe_file"]["sha256"], fresh_content_identity(self.odb)["sha256"])
        self.assertEqual(provenance["fe_geometry_identity"]["sha256"],
                         modal_dataset_geometry_identity(self.fe)["sha256"])
        rows = dict(sheet_rows(self.xlsx, "Provenance")[1:])
        self.assertEqual(rows["experimental_file.sha256"], provenance["experimental_file"]["sha256"])
        self.assertEqual(rows["mode_pairs"], "NONE")


class FrfPeakTests(_Case):
    def test_dataset_58_peaks_are_provisional_and_out_of_range_peaks_are_marked(self):
        unv = F.write_frf_only(self.folder / "frf_only.uff")
        short = ModalDataset("FE", self.fe.source_path, [m for m in self.fe.sorted_modes() if m.frequency_hz < 100.0],
                             metadata=dict(self.fe.metadata))
        report, pdf, xlsx = self.written(self.loaded(unv, fe=short))
        self.assertTrue(report.experimental_rows)
        self.assertTrue(all(r["classification"].startswith("PROVISIONAL") for r in report.experimental_rows))
        above = [r for r in report.experimental_rows if r["frequency_hz"] > report.fe_range_hz[1]]
        self.assertTrue(above)
        self.assertTrue(all("above the highest extracted FE mode" in r["warnings"] for r in above))
        self.assertTrue(all("above the highest extracted FE mode" in r["range_note"] for r in report.proximity_rows
                            if r["experimental_hz"] > report.fe_range_hz[1]))
        text = pdf_text(pdf)
        self.assertIn("PROVISIONAL", text)
        self.assertIn("above the highest extracted FE mode", text)
        self.assertNotIn("FITTED MODE", text)
        experimental = sheet_rows(xlsx, "Experimental Data")
        self.assertTrue(all(str(row[2]).startswith("PROVISIONAL") for row in experimental[1:]))


class UnusableExperimentTests(_Case):
    def test_fe_only_report_without_experimental_data(self):
        report, pdf, xlsx = self.written(self.loaded(None))
        self.assertEqual(report.status, "EXPERIMENTAL_DATA_NOT_USABLE")
        self.assertEqual((report.experimental_rows, report.proximity_rows), ((), ()))
        text = pdf_text(pdf)
        self.assertIn(dr.FE_ONLY_BANNER, text)
        self.assertIn("no modal or FRF record in the file", text)
        self.assertIn("No frequency proximity rows", text)
        self.assertEqual(len(sheet_rows(xlsx, "FE Modes")) - 1, len(self.fe.sorted_modes()))


class WriteFailureTests(_Case):
    def test_a_failed_write_leaves_no_partial_file_and_keeps_the_target(self):
        report = dr.build_diagnostic_report(self.loaded(S.write_experiment(self.folder / "f.unv", S.E0, S.G0)))
        target = self.folder / "locked.pdf"
        target.write_bytes(b"user file")
        with mock.patch("os.replace", side_effect=PermissionError("in use")):
            with self.assertRaises(dr.ReportWriteError) as caught:
                dr.write_diagnostic_pdf(report, target)
        self.assertIn("open in another program", str(caught.exception))
        self.assertEqual(target.read_bytes(), b"user file")
        self.assertEqual([p.name for p in self.folder.glob(".~report_*")], [])


class _Button:
    def __init__(self):
        self.options = {"state": "disabled", "text": ""}

    def configure(self, **options):
        self.options.update(options)

    def cget(self, key):
        return self.options[key]


class GuiTests(_Case):
    def app(self, loaded):
        messages = []
        app = SimpleNamespace(excel_button=_Button(), pdf_button=_Button(), status=SimpleNamespace(set=messages.append),
                              _post_to_ui=lambda callback: callback() or True, root=None, mode_viewer=None,
                              _diagnostic_export_source=None)
        import diagnostic_report_ui as ui

        ui.enable_diagnostic_export(app, loaded)
        return app, ui, messages

    def run_export(self, app, ui, kind, filename, **patches):
        shown = []
        with mock.patch.object(ui.filedialog, "asksaveasfilename", return_value=filename), \
                mock.patch.object(ui.messagebox, "showinfo", side_effect=lambda *a: shown.append(("info",) + a)), \
                mock.patch.object(ui.messagebox, "showerror", side_effect=lambda *a: shown.append(("error",) + a)), \
                mock.patch.object(ui.messagebox, "askyesno", return_value=patches.get("replace", False)), \
                mock.patch.object(ui.threading, "Thread",
                                  side_effect=lambda target, daemon: SimpleNamespace(start=target)):
            ui.export_report(app, kind)
        return shown

    def test_buttons_follow_the_result_state(self):
        loaded = self.loaded(S.write_experiment(self.folder / "f.unv", S.E0, S.G0))
        app, ui, _ = self.app(loaded)
        self.assertEqual((app.pdf_button.cget("state"), app.pdf_button.cget("text")),
                         ("normal", "Export diagnostic PDF"))
        ui.disable_report_export(app)
        self.assertEqual((app.pdf_button.cget("state"), app.pdf_button.cget("text")), ("disabled", "Export PDF report"))
        self.assertIsNone(app._diagnostic_export_source)
        ui.enable_diagnostic_export(app, loaded)
        ui.restore_comparison_export(app)
        self.assertEqual(app.excel_button.cget("text"), "Export Excel report")

    def test_cancel_wrong_extension_and_existing_file(self):
        loaded = self.loaded(S.write_experiment(self.folder / "f.unv", S.E0, S.G0))
        app, ui, _ = self.app(loaded)
        self.assertEqual(self.run_export(app, ui, "pdf", ""), [])  # cancelled: nothing written, nothing shown
        self.assertEqual(list(self.folder.glob("*.pdf")), [])
        shown = self.run_export(app, ui, "xlsx", str(self.folder / "report.txt"))
        self.assertTrue((self.folder / "report.txt.xlsx").exists())  # never a workbook named .txt
        self.assertEqual(shown[-1][0], "info")
        (self.folder / "again.txt.pdf").write_bytes(b"user file")
        self.assertEqual(self.run_export(app, ui, "pdf", str(self.folder / "again.txt"), replace=False), [])
        self.assertEqual((self.folder / "again.txt.pdf").read_bytes(), b"user file")  # not replaced without consent

    def test_write_error_is_shown_and_repeat_export_works(self):
        loaded = self.loaded(S.write_experiment(self.folder / "f.unv", S.E0, S.G0))
        app, ui, _ = self.app(loaded)
        with mock.patch("os.replace", side_effect=PermissionError("in use")):
            shown = self.run_export(app, ui, "pdf", str(self.folder / "r.pdf"))
        self.assertEqual(shown[-1][0], "error")
        self.assertFalse(app._diagnostic_export_running)
        shown = self.run_export(app, ui, "pdf", str(self.folder / "r.pdf"))
        self.assertEqual(shown[-1][0], "info")
        self.assertIn(dr.BANNER, pdf_text(self.folder / "r.pdf"))

    def test_the_report_describes_the_analysis_shown_when_the_export_started(self):
        first = self.loaded(S.write_experiment(self.folder / "f.unv", S.E0, S.G0))
        app, ui, _ = self.app(first)
        real = dr.build_diagnostic_report

        def build_then_new_analysis(loaded, **kwargs):
            report = real(loaded, **kwargs)
            ui.disable_report_export(app)  # a new analysis starts while the report is being written
            return report

        with mock.patch.object(ui.report_service, "build_diagnostic_report", side_effect=build_then_new_analysis):
            shown = self.run_export(app, ui, "xlsx", str(self.folder / "snap.xlsx"))
        self.assertIn("A new analysis was started meanwhile", shown[-1][2])
        rows = dict(sheet_rows(self.folder / "snap.xlsx", "Provenance")[1:])
        self.assertEqual(rows["experimental_file.sha256"], fresh_content_identity(self.folder / "f.unv")["sha256"])

    def test_comparison_export_is_unchanged_when_a_result_exists(self):
        import main
        import diagnostic_report_ui as ui

        application = SimpleNamespace(result=object(), cache=None)
        with mock.patch.object(ui, "export_report", side_effect=AssertionError("diagnostic export not expected")):
            main.app.ModalComparatorApp._pdf(application)  # the original handler (returns: no cache)
            main.app.ModalComparatorApp._excel(application)
        application = SimpleNamespace(result=None, _diagnostic_export_source=None)
        with mock.patch.object(ui, "export_report") as diagnostic:
            main.app.ModalComparatorApp._pdf(application)
        diagnostic.assert_called_once_with(application, "pdf")


if __name__ == "__main__":
    unittest.main()
