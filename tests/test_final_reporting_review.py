from pathlib import Path
import sys
import tempfile
import unittest

import numpy as np
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

import final_reporting_review as target
from final_reporting_review import install_final_reporting_review
from test_core import ModalCoreTests


def _fake_export_excel(_result, output_path, _image_directory):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.create_sheet("Mode Comparison")
    workbook.save(output_path)
    return output_path


class FinalReportingReviewTests(unittest.TestCase):
    def setUp(self):
        # install_final_reporting_review() is a process-wide, idempotent
        # monkeypatch (like every other install_* in this codebase); isolate
        # each test's view of it instead of relying on real reporting.py output
        # or on whatever other test modules may already have installed.
        self._original_installed = target._INSTALLED
        self._original_export_excel = target.reporting.export_excel
        target._INSTALLED = False
        target.reporting.export_excel = _fake_export_excel

    def tearDown(self):
        target._INSTALLED = self._original_installed
        target.reporting.export_excel = self._original_export_excel

    def make_result(self):
        return ModalCoreTests().synthetic_result()

    def test_export_excel_final_writes_signed_and_absolute_frequency_summary(self):
        result = self.make_result()
        errors = [pair.frequency_error_percent for pair in result.pairs]
        mac_values = [pair.mac for pair in result.pairs if pair.mac is not None]
        self.assertTrue(mac_values, "fixture must include at least one calculable MAC")

        install_final_reporting_review()
        export_excel_final = target.reporting.export_excel
        self.assertIsNot(export_excel_final, _fake_export_excel)

        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            output = export_excel_final(
                result, directory / "report.xlsx", directory / "images"
            )
            workbook = load_workbook(output)
            try:
                summary = workbook["Summary"]
                self.assertEqual(summary["A10"].value, "Mean absolute frequency error")
                self.assertAlmostEqual(
                    summary["B10"].value, float(np.mean(np.abs(errors))) / 100.0
                )
                self.assertEqual(summary["A11"].value, "Mean signed frequency error")
                self.assertAlmostEqual(
                    summary["B11"].value, float(np.mean(errors)) / 100.0
                )
                # Mean MAC must survive this layer rather than being clobbered
                # by the frequency-error rows that used to reuse the same cells.
                self.assertEqual(summary["A12"].value, "Mean MAC")
                self.assertAlmostEqual(summary["B12"].value, float(np.mean(mac_values)))
                self.assertEqual(summary["A13"].value, "Frequency-error sign")
                self.assertEqual(
                    workbook["Mode Comparison"].cell(row=1, column=5).value,
                    "Signed frequency error, %",
                )
            finally:
                workbook.close()

    def test_install_is_idempotent_and_does_not_rewrap(self):
        install_final_reporting_review()
        wrapped_once = target.reporting.export_excel

        install_final_reporting_review()

        self.assertIs(target.reporting.export_excel, wrapped_once)


if __name__ == "__main__":
    unittest.main()
