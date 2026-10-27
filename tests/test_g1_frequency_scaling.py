"""G1 — two-file frequency-scaling diagnostics (diagnostic only; never a material property or identification verdict).

The diagnostic reads the matched pairs of the existing comparison result (pairing is not redone) and reports
f_exp / f_FE, its square and ln(f_exp / f_FE) with MAC, comparator status and quality warnings. Only MAC-supported pairs
without warnings enter the descriptive median / range (never a statistical uncertainty). No E, G or ν is derived.
"""

from __future__ import annotations

import copy
import math
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import load_workbook  # noqa: E402

from domain.identification_pairing_policy import STRICT_IDENTIFICATION_PAIRING  # noqa: E402
import frequency_scaling as fs  # noqa: E402
import test_core  # noqa: E402


def pair(fe=100.0, exp=110.0, mac=0.95, status="Excellent match", order_changed=False, coverage="accepted",
         fe_mode=7, exp_mode=1):
    return SimpleNamespace(abaqus_mode=fe_mode, experimental_mode=exp_mode, abaqus_frequency_hz=fe,
                           experimental_frequency_hz=exp, mac=mac, status=status, order_changed=order_changed,
                           coverage_status=coverage)


def synthetic_result():
    return test_core.ModalCoreTests("test_export_excel_base_writes_the_documented_sheets_and_values").synthetic_result()


class RatioTests(unittest.TestCase):
    def test_exact_numerical_ratios(self):
        row = fs.scaling_row(pair(100.0, 110.0))
        self.assertAlmostEqual(row.frequency_ratio, 1.1, places=15)
        self.assertAlmostEqual(row.squared_frequency_ratio, 1.21, places=15)
        self.assertAlmostEqual(row.log_frequency_difference, math.log(1.1), places=15)
        row = fs.scaling_row(pair(28.006, 23.572))
        self.assertAlmostEqual(row.frequency_ratio, 23.572 / 28.006, places=15)
        self.assertAlmostEqual(row.squared_frequency_ratio, (23.572 / 28.006) ** 2, places=15)
        self.assertAlmostEqual(row.log_frequency_difference, math.log(23.572 / 28.006), places=15)

    def test_direction_of_frequency_differences(self):
        higher, lower = fs.scaling_row(pair(100, 105)), fs.scaling_row(pair(100, 95))
        equal = fs.scaling_row(pair(100, 100))
        self.assertEqual(higher.direction, "experiment higher than FE")
        self.assertGreater(higher.log_frequency_difference, 0.0)
        self.assertEqual(lower.direction, "experiment lower than FE")
        self.assertLess(lower.log_frequency_difference, 0.0)
        self.assertEqual(equal.direction, "equal")
        self.assertEqual(equal.log_frequency_difference, 0.0)

    def test_missing_and_invalid_frequencies(self):
        for fe, exp in ((None, 100.0), (100.0, None), (float("nan"), 100.0), (100.0, float("inf")), (0.0, 100.0),
                        (100.0, -5.0), ("100", 100.0), (True, 100.0)):
            with self.subTest(fe=fe, exp=exp):
                row = fs.scaling_row(pair(fe, exp))
                self.assertEqual(row.category, fs.INVALID_FREQUENCY)
                self.assertIsNone(row.frequency_ratio)
                self.assertIsNone(row.squared_frequency_ratio)
                self.assertIsNone(row.log_frequency_difference)
                self.assertFalse(row.in_summary)
        diagnostic = fs.frequency_scaling_diagnostic([pair(None, 100.0)])
        self.assertEqual(diagnostic.status, "UNAVAILABLE")
        self.assertIsNone(diagnostic.summary)

    def test_missing_pairs(self):
        for pairs in ([], None):
            diagnostic = fs.frequency_scaling_diagnostic(pairs)
            self.assertEqual(diagnostic.status, "UNAVAILABLE")
            self.assertEqual(diagnostic.rows, ())
            self.assertIn("no matched pairs", diagnostic.message)


class QualityTests(unittest.TestCase):
    def test_policy_threshold_is_the_accepted_strict_pairing_minimum(self):
        self.assertEqual(fs.MINIMUM_MAC, STRICT_IDENTIFICATION_PAIRING.minimum_mac)
        self.assertEqual(fs.MAXIMUM_RELATIVE_ERROR, STRICT_IDENTIFICATION_PAIRING.maximum_relative_frequency_error)

    def test_low_mac_frequency_only_and_questionable_matches(self):
        cases = ((pair(mac=None, status="Frequency match"), fs.FREQUENCY_ONLY, "frequency-only"),
                 (pair(mac=0.60, status="Review"), fs.LOW_MAC, "below the strict"),
                 (pair(mac=0.7999), fs.LOW_MAC, "below the strict"),
                 (pair(order_changed=True), fs.QUESTIONABLE, "order changed"),
                 (pair(coverage="insufficient point coverage"), fs.QUESTIONABLE, "coverage"),
                 (pair(100.0, 80.0, status="Review"), fs.QUESTIONABLE, "strict pairing window"),
                 (pair(mac=0.80), fs.MAC_SUPPORTED, None))
        for item, category, text in cases:
            with self.subTest(category=category, text=text):
                row = fs.scaling_row(item)
                self.assertEqual(row.category, category)
                self.assertEqual(row.in_summary, category == fs.MAC_SUPPORTED)
                if text:
                    self.assertIn(text, "; ".join(row.warnings))
                self.assertIsNotNone(row.frequency_ratio)  # values are shown, only flagged
        diagnostic = fs.frequency_scaling_diagnostic([c[0] for c in cases[:-1]])
        self.assertEqual(diagnostic.status, "UNAVAILABLE")
        self.assertIn("not evidence for material identification", diagnostic.message)
        self.assertEqual(len(diagnostic.rows), 6)

    def test_descriptive_median_and_spread(self):
        pairs = [pair(100.0, 101.0), pair(200.0, 196.0, fe_mode=8), pair(300.0, 309.0, fe_mode=9),
                 pair(100.0, 140.0, mac=None, fe_mode=10)]  # the frequency-only pair is shown, not summarised
        diagnostic = fs.frequency_scaling_diagnostic(pairs)
        self.assertEqual(diagnostic.status, "AVAILABLE")
        summary = diagnostic.summary
        self.assertEqual(summary.count, 3)
        self.assertAlmostEqual(summary.median_ratio, 1.01, places=15)
        self.assertAlmostEqual(summary.median_squared_ratio, 1.01 ** 2, places=15)
        self.assertAlmostEqual(summary.median_log_difference, math.log(1.01), places=15)
        self.assertAlmostEqual(summary.min_log_difference, math.log(0.98), places=15)
        self.assertAlmostEqual(summary.max_log_difference, math.log(1.03), places=15)
        self.assertIn("not a statistical uncertainty", diagnostic.message)
        single = fs.frequency_scaling_diagnostic([pair(100.0, 101.0)])
        self.assertIn("no spread", single.message)
        self.assertEqual(single.summary.count, 1)

    def test_labels_and_no_material_values(self):
        self.assertEqual(fs.LABELS, ("FREQUENCY_SCALING_DIAGNOSTIC_ONLY", "NOT_A_MATERIAL_PROPERTY",
                                     "NOT_A_SCIENTIFIC_IDENTIFICATION_VERDICT"))
        source = (ROOT / "src" / "frequency_scaling.py").read_text(encoding="utf-8")
        for forbidden in ("youngs", "poisson", "shear_modulus", "E_in_plane", "G12", "sensitivity",
                          "campaign", "journal", "judge_campaign_run", "subprocess"):
            self.assertNotIn(forbidden, source)
        row = fs.scaling_row(pair())
        self.assertFalse(hasattr(row, "youngs_modulus"))


class IntegrationTests(unittest.TestCase):
    def test_comparison_pairs_are_read_not_changed(self):
        result = synthetic_result()
        before = [copy.deepcopy(p.__dict__) for p in result.pairs]
        diagnostic = fs.frequency_scaling_diagnostic(result.pairs)
        after = [p.__dict__ for p in result.pairs]
        self.assertEqual(len(before), len(after))
        for old, new in zip(before, after):
            self.assertEqual(old.keys(), new.keys())
            for key in old:
                if hasattr(old[key], "shape"):
                    self.assertTrue((old[key] == new[key]).all())
                else:
                    self.assertEqual(old[key], new[key])
        self.assertEqual([(r.fe_mode, r.experimental_mode) for r in diagnostic.rows],
                         [(p.abaqus_mode, p.experimental_mode) for p in result.pairs])
        for row, item in zip(diagnostic.rows, result.pairs):
            self.assertAlmostEqual(row.frequency_ratio, item.experimental_frequency_hz / item.abaqus_frequency_hz)

    def test_excel_sheet(self):
        import reporting

        result = synthetic_result()
        with tempfile.TemporaryDirectory() as directory:
            path = reporting.export_excel(result, Path(directory) / "report.xlsx", Path(directory) / "images")
            workbook = load_workbook(path)
            self.assertIn("Frequency Scaling", workbook.sheetnames)
            sheet = workbook["Frequency Scaling"]
            values = [[cell.value for cell in row] for row in sheet.iter_rows()]
            self.assertEqual(values[0][:2], ["Labels", fs.LABEL_TEXT])
            self.assertEqual(values[1][0], "Caveat")
            header = values.index([*fs.EXCEL_HEADERS])
            rows = values[header + 1:]
            self.assertEqual(len(rows), len(result.pairs))
            for data, item in zip(rows, result.pairs):
                self.assertAlmostEqual(data[5], item.experimental_frequency_hz / item.abaqus_frequency_hz)
                self.assertAlmostEqual(data[6], (item.experimental_frequency_hz / item.abaqus_frequency_hz) ** 2)
                self.assertAlmostEqual(data[7], math.log(item.experimental_frequency_hz / item.abaqus_frequency_hz))
            comparison = workbook["Mode Comparison"]  # the existing sheet is unchanged
            self.assertEqual(comparison.max_row, len(result.pairs) + 1)

    def test_excel_unavailable_without_pairs(self):
        from openpyxl import Workbook

        workbook = Workbook()
        fs.write_excel_sheet(workbook, fs.frequency_scaling_diagnostic([]))
        values = [[cell.value for cell in row] for row in workbook["Frequency Scaling"].iter_rows()]
        self.assertEqual(values[2][:2], ["Status", "UNAVAILABLE"])

    def test_real_tk_comparison_table(self):
        try:
            import tkinter as tk

            root = tk.Tk()
        except Exception as error:  # noqa: BLE001
            self.skipTest(f"Tk is unavailable: {error}")
        self.addCleanup(root.destroy)
        root.withdraw()
        import main
        import ui_workflow

        disabled = {"autosave_enabled": False, "restore_on_startup": False, "show_recovery_notice": False}
        with mock.patch.object(ui_workflow, "load_recovery_preferences", return_value=disabled), \
                mock.patch.object(ui_workflow, "discover_abaqus_installations", return_value=()):
            application = main.app.ModalComparatorApp(root)
            application._show_expert()  # APP-UI.1: the window opens on Home; these are expert-tab checks
        self.assertIn("no comparison result yet", application.scaling_summary.cget("text"))
        result = synthetic_result()
        with tempfile.TemporaryDirectory() as directory:
            application.cache = Path(directory)
            application.result = result
            application._populate(result)
            root.update_idletasks()
        table = application.scaling_table
        items = [tuple(str(v) for v in table.item(k, "values")) for k in table.get_children()]
        self.assertEqual(len(items), len(result.pairs))
        expected = fs.table_rows(fs.frequency_scaling_diagnostic(result.pairs))
        self.assertEqual(items, [tuple(str(v) for v in row) for row in expected])
        self.assertTrue(application.scaling_summary.cget("text").startswith(("AVAILABLE", "UNAVAILABLE")))
        self.assertEqual(len(application.table.get_children()), len(result.pairs))  # comparison table unchanged

    def test_real_tk_scrolling_keeps_every_column_accessible(self):
        # PR #71 review: all 12 columns, especially the last three, stay reachable at the minimum window widths
        try:
            import tkinter as tk

            root = tk.Tk()
        except Exception as error:  # noqa: BLE001
            self.skipTest(f"Tk is unavailable: {error}")
        self.addCleanup(root.destroy)
        # runs before destroy: cancel pending idle callbacks of the shown window
        self.addCleanup(lambda: [root.after_cancel(job) for job in root.tk.splitlist(root.tk.call("after", "info"))])
        import main
        import ui_workflow
        from ui_policy import MINIMUM_WINDOW_SIZE

        disabled = {"autosave_enabled": False, "restore_on_startup": False, "show_recovery_notice": False}
        with mock.patch.object(ui_workflow, "load_recovery_preferences", return_value=disabled), \
                mock.patch.object(ui_workflow, "discover_abaqus_installations", return_value=()):
            application = main.app.ModalComparatorApp(root)
            application._show_expert()  # APP-UI.1: the window opens on Home; these are expert-tab checks
        table = application.scaling_table
        self.assertEqual(str(application.scaling_xscroll.cget("orient")), "horizontal")
        self.assertEqual(str(application.scaling_yscroll.cget("orient")), "vertical")
        self.assertTrue(str(table.cget("xscrollcommand")) and str(table.cget("yscrollcommand")))
        rows = [pair(100.0 + i, 101.0 + i, fe_mode=7 + i, exp_mode=1 + i,
                     order_changed=True, coverage="insufficient spatial coverage") for i in range(20)]
        for row in fs.table_rows(fs.frequency_scaling_diagnostic(rows)):
            table.insert("", "end", values=row)
        # its own tab, appended last; the comparison-table tab gains no widget (its layout is unchanged)
        tabs = application.tabs.tabs()
        self.assertEqual(str(tabs[-1]), str(application.scaling_tab))
        from ui_policy import COMPACT_TAB_LABELS, FULL_TAB_LABELS

        self.assertEqual(main.app.SCALING_TAB_TITLE, FULL_TAB_LABELS[-1])
        self.assertIn(application.tabs.tab(application.scaling_tab, "text"),
                      (FULL_TAB_LABELS[-1], COMPACT_TAB_LABELS[-1]))  # responsive full / compact label
        self.assertLess(tabs.index(str(application.table_tab)), tabs.index(str(application.details_tab)))
        self.assertNotIn(str(application.scaling_tab), [str(w) for w in application.table_tab.winfo_children()])
        self.assertFalse(any(str(w).startswith(str(application.table_tab)) for w in
                             (table, application.scaling_summary)))
        application.tabs.select(application.scaling_tab)
        columns = table.cget("columns")
        self.assertEqual(len(columns), 12)
        for width, height in ((1160, 720), MINIMUM_WINDOW_SIZE):  # app.py minsize and the responsive minimum
            with self.subTest(width=width, height=height):
                root.geometry(f"{width}x{height}+0+0")
                root.deiconify()
                root.update()
                table.xview_moveto(0.0)
                table.yview_moveto(0.0)
                root.update()
                self.assertTrue(table.winfo_ismapped())  # fully shown at the minimum window size
                self.assertTrue(application.scaling_xscroll.winfo_ismapped())
                self.assertTrue(application.scaling_yscroll.winfo_ismapped())
                self.assertLess(table.xview()[1], 1.0)  # wider than the window: scrolling is required
                for key in columns:  # every column is fixed (never squeezed below its width)
                    self.assertFalse(int(table.column(key, "stretch")))
                table.xview_moveto(1.0)
                root.update()
                self.assertEqual(table.xview()[1], 1.0)
                first = table.get_children()[0]
                visible = table.winfo_width()
                for key in ("category", "status", "warnings"):
                    box = table.bbox(first, key)
                    self.assertTrue(box, key)  # reachable through horizontal scrolling
                    self.assertLessEqual(box[0] + box[2], visible + 1, key)
                self.assertLess(table.yview()[1], 1.0)  # more rows than the table height
                table.yview_moveto(1.0)
                root.update()
                self.assertTrue(table.bbox(table.get_children()[-1], "warnings"))


if __name__ == "__main__":
    unittest.main()
