"""NEW-GUI.6: the verified modal comparison, its report, the material evidence and the known sources (no Qt).

``new_gui.comparison_model`` verifies the archived registration again with the existing service and runs
``modal_core.compare_modal_datasets(..., frozen_registration=...)`` on the data Inspect file read. On the project's
synthetic fixture world (G2c / G2d-1):

- the verified comparison is the comparator's own result (pairs and MAC matrix equal to a direct call), with no
  automatic orientation search and with the registration's provenance attached as the existing program attaches it;
- a changed experimental file (also after a successful check), another FE geometry, another modal set, another
  registration or an FRF set (dataset 58, PROVISIONAL) is refused before the comparator runs;
- zero admissible pairs is a safe result without an invented correlation; close modes and order changes are shown as
  the comparator reports them; experimental record 1 stays visible; rigid-body FE modes are separated by the
  comparator's shape criterion and the first elastic mode is kept;
- a cancelled request never runs the comparator, a late cancellation discards the result, and the comparison shares
  the one heavy-operation lock of NEW-GUI.4/5;
- the pair drawings come from the program's final owner of ``reporting.render_pair_images`` and leave no file;
- the report is written only into the chosen folder: the verified report keeps the statuses, the diagnostic report
  has no MAC, a changed file is refused; the program's report layers are installed in src/main.py's order;
- the material records are read-only and classified; no material value appears in the material-identification class,
  and nothing is related to data without a verified fixture;
- the known sources are confirmed by a typed reading of the provenance map, never by a substring of its text.

With ``NEW_GUI6_REAL_INPUTS`` set to a folder holding verified read-only copies (``SP02/fe_extraction/manifest.json``,
``SP02/SP02_polymax_retry_260803.unv``, ``SP13/fe_extraction/manifest.json``, ``SP13/SP13_a_polymax.unv``), the
positive path runs on the real archived SP-02 and SP-13 data; otherwise those tests are skipped.
"""

from __future__ import annotations

import ast
import contextlib
import csv
import importlib
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import threading
import unittest
from unittest import mock

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

import modal_core  # noqa: E402
import reviewed_core  # noqa: E402
from modal_core import ModalDataset, ModeShape  # noqa: E402
from new_gui import comparison_model as cm  # noqa: E402
from new_gui import diagnostic_model as dm  # noqa: E402
from new_gui import known_sources as ks  # noqa: E402
from new_gui import materials_model as mm  # noqa: E402
from new_gui import preview_model as pm  # noqa: E402
from new_gui import registration_check_model as rm  # noqa: E402
from new_gui import report_export as rx  # noqa: E402
from services import new_odb_comparison as service  # noqa: E402
import frf_unv_support as F  # noqa: E402
from test_g2d1_governed_registration import experimental_dataset, fe_dataset  # noqa: E402
from test_new_gui5_registration import WorldCase, experimental_preview, fe_preview, fingerprint, other_set  # noqa: E402
from test_new_odb_comparison import new_odb_dataset  # noqa: E402

REAL = os.environ.get("NEW_GUI6_REAL_INPUTS", "")


def with_frequencies(dataset: ModalDataset, frequencies: dict) -> ModalDataset:
    """The same FE mesh and shapes with other frequencies (the FE geometry identity is unchanged)."""
    modes = [ModeShape(m.number, frequencies.get(m.number, m.frequency_hz), m.node_ids, m.coordinates, m.vectors)
             for m in dataset.sorted_modes()]
    return ModalDataset(dataset.source_name, dataset.source_path, modes, metadata=dict(dataset.metadata))


def with_rigid_mode(dataset: ModalDataset) -> ModalDataset:
    """A rigid-body translation mode added below the elastic modes, on the same mesh."""
    first = dataset.sorted_modes()[0]
    rigid = ModeShape(1, 0.01, first.node_ids, first.coordinates, np.tile([0.0, 0.0, 1.0], (len(first.node_ids), 1)))
    return ModalDataset(dataset.source_name, dataset.source_path, [rigid] + dataset.sorted_modes(),
                        metadata=dict(dataset.metadata))


def dataset_state(dataset: ModalDataset) -> tuple:
    """Everything of a dataset a comparison could change."""
    return (json.dumps(dataset.metadata, sort_keys=True, default=str),
            tuple((m.number, m.frequency_hz, m.vectors.tobytes(), m.coordinates.tobytes(),
                   json.dumps(m.metadata, sort_keys=True, default=str)) for m in dataset.modes))


class _Comparison(WorldCase):
    def setUp(self):
        super().setUp()
        self.fe_preview, self.exp_preview = self.previews()
        self.registration = self.world.registrations["SP02"].registration_hash
        own = tempfile.TemporaryDirectory()
        self.addCleanup(own.cleanup)
        self.own = Path(own.name)
        # NEW-GUI.6-F1: the identities Inspect file records when it reads the files
        self.sources = {"fe": pm.file_identity(self.odb), "experimental": pm.file_identity(self.unv)}

    @contextlib.contextmanager
    def changed_unv(self, content: bytes):
        """The experimental file with other content, restored before the world checks that nothing changed."""
        with self.changed(self.unv, content):
            yield

    @contextlib.contextmanager
    def changed(self, path: Path, content):
        """A world file with other content, restored before the world checks that nothing changed."""
        original = path.read_bytes()
        path.write_bytes(content(original))
        try:
            yield
        finally:
            path.write_bytes(original)

    def previews(self, fe=None, source=None, datasets=None):
        source = source or self.unv
        fe = fe or self.fe
        datasets = datasets if datasets is not None else [experimental_dataset(source, self.fe),
                                                          other_set(source, self.fe)]
        return fe_preview(self.odb, fe), experimental_preview(source, datasets)

    def run_compare(self, fe_prev=None, exp_prev=None, set_index=0, registration="default", cancel=None,
                    sources="default"):
        registration = self.registration if registration == "default" else registration
        sources = self.sources if sources == "default" else sources
        return cm.run_compare(fe_prev or self.fe_preview, exp_prev or self.exp_preview, set_index, registration, 1,
                              cancel or threading.Event(), sources)


class VerifiedComparisonTests(_Comparison):
    def test_the_result_is_the_comparators_own_with_the_frozen_registration(self):
        real = modal_core.compare_modal_datasets
        calls = []

        def spy(*args, **kwargs):
            calls.append(kwargs)
            return real(*args, **kwargs)

        bind = mock.Mock(side_effect=service.bind_registration_by_fe_geometry)
        with mock.patch.object(modal_core, "compare_modal_datasets", side_effect=spy), \
                mock.patch.object(service, "bind_registration_by_fe_geometry", bind), \
                mock.patch.object(reviewed_core, "geometry_alignment_candidates",
                                  side_effect=AssertionError("no automatic orientation search")):
            outcome = self.run_compare()
        self.assertEqual(outcome.error, "")
        self.assertEqual(outcome.refusal, "")
        self.assertEqual(len(calls), 1)
        self.assertEqual(set(calls[0]), {"frozen_registration"})  # nothing else: no scale, orientation or mapping
        self.assertEqual(calls[0]["frozen_registration"].registration_hash, self.registration)
        self.assertEqual(bind.call_count, 1)  # verified again at run time, not taken from a shown status
        self.assertEqual(Path(bind.call_args.args[0]), self.unv)
        direct = real(self.fe, experimental_dataset(self.unv, self.fe),
                      frozen_registration=self.world.registrations["SP02"])
        result = outcome.result
        self.assertIsInstance(result, modal_core.ComparisonResult)
        self.assertEqual([(p.abaqus_mode, p.experimental_mode) for p in result.pairs],
                         [(p.abaqus_mode, p.experimental_mode) for p in direct.pairs])
        self.assertEqual([(8, 2), (10, 3), (13, 4)], [(p.abaqus_mode, p.experimental_mode) for p in result.pairs])
        np.testing.assert_array_equal(np.nan_to_num(result.mac_matrix, nan=-1),
                                      np.nan_to_num(direct.mac_matrix, nan=-1))
        self.assertEqual([p.mac for p in result.pairs], [p.mac for p in direct.pairs])
        self.assertEqual(result.metadata["orientation_source"], reviewed_core.FROZEN_REGISTRATION_ORIENTATION_SOURCE)
        self.assertEqual(result.metadata["geometry_verified_registration"]["registration_hash"], self.registration)
        self.assertTrue(result.warnings[0].startswith("Physical registration reused by FE geometry identity"))
        view = outcome.view
        self.assertEqual(view.status, cm.PAIRS_FOUND)
        self.assertEqual(view.fixture_id, "SP02/bravo-1-physical")
        self.assertEqual([r.cells[6] for r in view.pair_rows], [f"{p.mac:.3f}" for p in result.pairs])
        self.assertEqual([r.cells[5] for r in view.pair_rows],
                         [f"{p.frequency_error_percent:+.2f}" for p in result.pairs])
        self.assertEqual({r.cells[13] for r in view.pair_rows}, {cm.AUTO_ID_USE})
        self.assertEqual(len(view.mac_cells), len(result.abaqus_mode_numbers))
        self.assertEqual(len(view.assigned_cells), 3)
        self.assertTrue(outcome.mac_image.startswith(b"\x89PNG"))

    def test_the_read_data_are_not_changed_and_nothing_is_written(self):
        before = (dataset_state(self.fe_preview.dataset), dataset_state(self.exp_preview.sets[0].dataset))
        temp = Path(tempfile.gettempdir())
        leftovers = {p.name for p in temp.glob("new_gui_*")}
        outcome = self.run_compare()
        self.assertIsNotNone(outcome.view)
        self.assertEqual((dataset_state(self.fe_preview.dataset), dataset_state(self.exp_preview.sets[0].dataset)),
                         before)
        self.assertEqual(fingerprint(self.world.root), self.before)
        self.assertEqual({p.name for p in temp.glob("new_gui_*")}, leftovers)  # the MAC drawing folder is removed

    def test_experimental_record_one_stays_visible_and_unpaired_modes_are_listed(self):
        view = self.run_compare().view
        candidates = {row[1]: row for row in view.candidate_rows}
        self.assertEqual(sorted(candidates), ["1", "2", "3", "4"])  # every experimental mode, record 1 included
        self.assertEqual(candidates["1"][0], "1")
        self.assertEqual(candidates["1"][10], "no")
        self.assertIn("frequency", candidates["1"][9])  # 24 Hz vs FE 7 at 20 Hz: outside the 15 % gate
        self.assertTrue(any(text.startswith("Exp record 1 (mode 1") for text in view.unpaired_experimental))
        self.assertIn(1, view.mac_experimental_modes)

    def test_rigid_body_fe_modes_are_separated_by_shape_and_the_first_elastic_mode_is_kept(self):
        fe = with_rigid_mode(self.fe)
        fe_prev, exp_prev = self.previews(fe=fe)
        view = self.run_compare(fe_prev, exp_prev).view
        self.assertEqual([row[0] for row in view.rigid_rows], ["1"])
        self.assertEqual(view.mac_fe_modes[0], 7)  # the first elastic FE mode is compared
        self.assertNotIn(1, view.mac_fe_modes)
        self.assertTrue(any("rigid-body mode(s) by shape projection" in w for w in view.warnings))

    def test_zero_admissible_pairs_is_a_safe_result_without_correlation(self):
        fe = with_frequencies(self.fe, {n: 3.0 * hz for n, hz in
                                        ((m.number, m.frequency_hz) for m in self.fe.sorted_modes())})
        fe_prev, exp_prev = self.previews(fe=fe)
        outcome = self.run_compare(fe_prev, exp_prev)
        self.assertEqual(outcome.error, "")
        self.assertEqual(outcome.result.pairs, [])
        self.assertEqual(outcome.view.status, cm.NO_PAIR)
        self.assertIn("valid scientific result", outcome.view.statement)
        self.assertEqual(outcome.view.pair_rows, ())
        self.assertEqual(outcome.view.assigned_cells, frozenset())
        self.assertIn("No pair", outcome.view.target_counts)
        self.assertEqual({row[10] for row in outcome.view.candidate_rows}, {"no"})

    def test_close_modes_and_order_changes_are_shown_as_the_comparator_reports_them(self):
        fe = with_frequencies(self.fe, {8: 80.0, 9: 120.0, 10: 78.0})
        fe_prev, exp_prev = self.previews(fe=fe)
        outcome = self.run_compare(fe_prev, exp_prev)
        result, view = outcome.result, outcome.view
        flagged = {(p.abaqus_mode, p.experimental_mode): p.order_changed for p in result.pairs}
        self.assertTrue(any(flagged.values()))
        for row in view.pair_rows:
            pair = result.pairs[row.index]
            self.assertEqual(row.cells[11], "changed" if pair.order_changed else "kept")
        groups = result.abaqus.metadata["close_abaqus_mode_groups"]
        self.assertTrue(groups)
        self.assertTrue(any("FE modes" in text for text in view.close_groups))
        self.assertTrue(any("Closely spaced Abaqus modes" in w for w in view.warnings))
        closes = [row.cells[12] for row in view.pair_rows if row.abaqus_mode in (8, 10)]
        self.assertTrue(all(text.startswith("FE ") for text in closes))

    def test_a_pair_without_mac_is_never_shown_as_a_shape_identity(self):
        outcome = self.run_compare()
        pair = outcome.result.pairs[0]
        pair.mac = None  # a frequency-only assignment, as the comparator may make one
        view = cm.comparison_view(outcome.result, outcome.bound, self.fe_preview, self.exp_preview, 0)
        self.assertIn(cm.MAC_NOT_COMPUTED, view.pair_rows[0].cells[6])
        self.assertIn("shape identity NOT established", view.pair_rows[0].cells[6])


class RefusalTests(_Comparison):
    def assert_refused_before_the_comparator(self, outcome, text):
        self.assertIsNone(outcome.view)
        self.assertIn(text, outcome.refusal)

    def compare_without_comparator(self, **kwargs):
        with mock.patch.object(modal_core, "compare_modal_datasets",
                               side_effect=AssertionError("the comparator must not run")):
            return self.run_compare(**kwargs)

    def test_a_changed_experimental_file_is_refused_before_mac(self):
        with self.changed_unv(lambda data: data + b"x"):
            outcome = self.compare_without_comparator()
        self.assert_refused_before_the_comparator(outcome,
                                                  "is not the PolyMAX source of an accepted physical registration")

    def test_a_file_changed_after_a_successful_check_is_refused_at_run_time(self):
        check = rm.run_check(self.fe_preview, self.exp_preview, 0, 1, threading.Event())
        self.assertIsNotNone(check.view)
        with self.changed_unv(lambda data: data[:-1] + b"#"):  # same size, other content
            outcome = self.compare_without_comparator()
        self.assert_refused_before_the_comparator(outcome, "SHA-256")

    def test_another_fe_geometry_is_refused(self):
        fe_prev, exp_prev = self.previews(fe=new_odb_dataset(self.odb, self.fe, shift=0.5))
        self.assert_refused_before_the_comparator(
            self.compare_without_comparator(fe_prev=fe_prev, exp_prev=exp_prev), "is not the accepted SP02 FE")

    def test_another_modal_set_is_refused(self):
        self.assert_refused_before_the_comparator(self.compare_without_comparator(set_index=1), "set-b")

    def test_an_fe_manifest_changed_since_it_was_read_is_refused_before_the_comparator(self):
        with self.changed(self.odb, lambda data: data + b"#"):
            outcome = self.compare_without_comparator()
        self.assert_refused_before_the_comparator(outcome, "is not the file that was read")

    def test_an_experimental_file_other_than_the_one_read_is_refused(self):
        # the read data came from other bytes than the (accepted) file on disk now
        sources = dict(self.sources, experimental=dict(self.sources["experimental"], sha256="0" * 64))
        self.assert_refused_before_the_comparator(self.compare_without_comparator(sources=sources),
                                                  "is not the file that was read")

    def test_the_outcome_keeps_the_identities_recorded_at_the_read(self):
        outcome = self.run_compare()
        self.assertEqual(outcome.sources, self.sources)

    def test_another_registration_than_the_checked_one_is_refused(self):
        self.assert_refused_before_the_comparator(self.compare_without_comparator(registration="0" * 64),
                                                  "check the registration again")

    def test_frf_records_are_provisional_and_never_compared(self):
        frf = F.write_frf_only(self.own / "frf_only.uff")
        from new_gui import preview_model as pm
        from new_gui import workspace_state as ws

        read = pm.run_read(ws.EXPERIMENTAL, frf, 1, threading.Event())
        self.assertIsNotNone(read.preview, read.error)
        self.assertTrue(rm.set_is_provisional(read.preview, 0))
        self.assertIn(cm.PROVISIONAL_BLOCK, cm.unavailable(self.fe_preview, read.preview, 0, object()))
        with mock.patch.object(service, "bind_registration_by_fe_geometry",
                               side_effect=AssertionError("no registration check for FRF records")):
            outcome = self.compare_without_comparator(exp_prev=read.preview)
        self.assert_refused_before_the_comparator(outcome, "PROVISIONAL — NOT A VALIDATED FITTED MODE")

    def test_without_a_verified_registration_the_comparison_is_unavailable(self):
        self.assertEqual(cm.unavailable(None, None, None, None), (dm.NEED_FE, dm.NEED_EXPERIMENTAL))
        self.assertEqual(cm.unavailable(self.fe_preview, self.exp_preview, None, None), (dm.NEED_SET,))
        self.assertEqual(cm.unavailable(self.fe_preview, self.exp_preview, 0, None), (cm.NEED_REGISTRATION,))


class RequestTests(_Comparison):
    def test_a_cancelled_request_waiting_for_the_lock_never_runs_the_comparator(self):
        cancel = threading.Event()
        with mock.patch.object(modal_core, "compare_modal_datasets",
                               side_effect=AssertionError("the comparator must not run")):
            with rm.SERVICE_LOCK:  # a frequency diagnostic or check is running
                done = []
                worker = threading.Thread(target=lambda: done.append(self.run_compare(cancel=cancel)))
                worker.start()
                threading.Event().wait(0.2)
                self.assertFalse(done)  # it waits for the shared lock
                cancel.set()
            worker.join(30)
        self.assertTrue(done[0].cancelled)

    def test_one_heavy_operation_at_a_time_shared_with_the_diagnostic_and_the_check(self):
        self.assertIs(rm.SERVICE_LOCK, dm._SERVICE)
        inside = []
        real = modal_core.compare_modal_datasets

        def spy(*args, **kwargs):
            inside.append(rm.SERVICE_LOCK.locked())
            return real(*args, **kwargs)

        with mock.patch.object(modal_core, "compare_modal_datasets", side_effect=spy):
            self.assertIsNotNone(self.run_compare().view)
        self.assertEqual(inside, [True])

    def test_a_late_cancellation_discards_the_result(self):
        cancel = threading.Event()
        real = modal_core.compare_modal_datasets

        def cancel_during(*args, **kwargs):
            result = real(*args, **kwargs)
            cancel.set()
            return result

        with mock.patch.object(modal_core, "compare_modal_datasets", side_effect=cancel_during):
            outcome = self.run_compare(cancel=cancel)
        self.assertTrue(outcome.cancelled)
        self.assertIsNone(outcome.result)

    def test_only_the_current_request_is_shown(self):
        slot = cm.ComparisonSlot()
        token, cancel = slot.start("first")
        slot.reset()  # another set, a replaced file, Inspect again, a closed window
        self.assertTrue(cancel.is_set())
        self.assertFalse(slot.finish(cm.ComparisonOutcome(token, view=object())))
        token, _ = slot.start("second")
        self.assertTrue(slot.finish(cm.ComparisonOutcome(token, view=object())))
        self.assertEqual(slot.state, cm.DONE)


class PairImageTests(_Comparison):
    def test_the_pair_is_drawn_by_the_programs_final_owner_and_leaves_no_file(self):
        import amplitude_correlation

        result = self.run_compare().result
        temp = Path(tempfile.gettempdir())
        leftovers = {p.name for p in temp.glob("new_gui_pair_*")}
        with mock.patch.object(amplitude_correlation, "render_pair_images",
                               wraps=amplitude_correlation.render_pair_images) as owner:
            images = cm.run_pair_images(result, 1, 7, threading.Event())
        self.assertEqual(images.error, "")
        self.assertIs(owner.call_args.args[0], result.pairs[1])
        self.assertEqual([caption for caption, _ in images.images], [c for _, c in cm.PAIR_CAPTIONS])
        self.assertTrue(all(data.startswith(b"\x89PNG") for _, data in images.images))
        self.assertEqual({p.name for p in temp.glob("new_gui_pair_*")}, leftovers)

    def test_a_cancelled_drawing_is_not_shown(self):
        result = self.run_compare().result
        cancel = threading.Event()
        cancel.set()
        self.assertTrue(cm.run_pair_images(result, 0, 1, cancel).cancelled)


class FinalOwnerTests(unittest.TestCase):
    def test_the_final_owners_are_the_programs(self):
        code = ("import sys; sys.path.insert(0, sys.argv[1]); import main, reporting, amplitude_correlation, "
                "reporting_hardening; print(reporting.render_pair_images is amplitude_correlation.render_pair_images,"
                " reporting.render_mac_matrix is reporting_hardening.render_mac_matrix)")
        import subprocess

        done = subprocess.run([sys.executable, "-c", code, str(ROOT / "src")], capture_output=True, text=True,
                              timeout=240, cwd=str(ROOT), env={**os.environ, "PYTHONIOENCODING": "utf-8",
                                                               "PYTHONDONTWRITEBYTECODE": "1"})
        self.assertEqual(done.returncode, 0, done.stderr[-2000:])
        self.assertEqual(done.stdout.strip().splitlines()[-1], "True True")


@contextlib.contextmanager
def reporting_restored():
    """Restore the report functions and layer flags the program's report layers change (test isolation)."""
    import advanced_metrics
    import reporting

    modules = [importlib.import_module(name) for name, _ in rx.REPORT_LAYERS]
    saved = {module: {k: v for k, v in vars(module).items() if k.startswith(("_INSTALLED", "_ORIGINAL"))}
             for module in modules}
    functions = {k: v for k, v in vars(reporting).items() if callable(v)}
    common = advanced_metrics._common_pair_data
    try:
        yield
    finally:
        for module, values in saved.items():
            vars(module).update(values)
        vars(reporting).update(functions)
        advanced_metrics._common_pair_data = common


class ReportTests(_Comparison):
    def setUp(self):
        super().setUp()
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.folder = Path(folder.name)

    def test_the_report_layers_are_main_pys_in_its_order(self):
        tree = ast.parse((ROOT / "src" / "main.py").read_text(encoding="utf-8"))
        calls = [node.func.id for node in ast.walk(tree) if isinstance(node, ast.Call)
                 and isinstance(node.func, ast.Name) and node.func.id.startswith("install_")]
        names = [name for _, name in rx.REPORT_LAYERS]
        self.assertEqual([call for call in calls if call in names], names)
        imports = {alias.name: node.module for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)
                   for alias in node.names}
        self.assertEqual([(imports[name], name) for name in names], list(rx.REPORT_LAYERS))

    def test_the_verified_report_keeps_the_statuses_and_writes_only_into_the_folder(self):
        outcome = self.run_compare()
        before = fingerprint(self.world.root)
        with reporting_restored():
            written = rx.export_verified(outcome.result, outcome.bound, outcome.view, self.fe_preview,
                                         self.exp_preview, 0, self.folder, sources=outcome.sources)
        self.assertEqual(sorted(p.suffix for p in written), [".json", ".pdf", ".xlsx"])
        self.assertEqual(sorted(p.name for p in self.folder.iterdir()), sorted(p.name for p in written))
        self.assertEqual(fingerprint(self.world.root), before)
        record = json.loads(next(p for p in written if p.suffix == ".json").read_text(encoding="utf-8"))
        self.assertEqual(record["status"], cm.PAIRS_FOUND)
        self.assertEqual(record["registration"]["registration_hash"], self.registration)
        self.assertEqual(record["experimental_source"]["sha256"], outcome.bound.experimental["sha256"])
        self.assertEqual(record["fe_source"]["sha256"], self.sources["fe"]["sha256"])  # the identity of the read
        self.assertTrue(record["fe_source"]["unchanged_at_export"])
        self.assertEqual([(p["fe_mode"], p["experimental_mode"], p["mac"]) for p in record["pairs"]],
                         [(p.abaqus_mode, p.experimental_mode, p.mac) for p in outcome.result.pairs])
        self.assertEqual({p["auto_id_use"] for p in record["pairs"]}, {cm.AUTO_ID_USE})
        self.assertEqual(record["suitability"]["material_property"], "none identified or implied")
        from openpyxl import load_workbook

        workbook = load_workbook(next(p for p in written if p.suffix == ".xlsx"))
        self.assertIn("Mode Comparison", workbook.sheetnames)

    def export(self, outcome, sources="outcome"):
        return rx.export_verified(outcome.result, outcome.bound, outcome.view, self.fe_preview, self.exp_preview, 0,
                                  self.folder, sources=outcome.sources if sources == "outcome" else sources)

    def test_a_file_changed_since_the_comparison_is_refused(self):
        outcome = self.run_compare()
        for path in (self.unv, self.odb):
            with self.subTest(file=path.name), self.changed(path, lambda data: data + b"x"):
                with self.assertRaisesRegex(rx.ExportRefused, "changed since it was read"):
                    self.export(outcome)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_without_identities_recorded_at_the_read_no_provenance_is_claimed(self):
        outcome = self.run_compare()
        for sources in (None, {"fe": None, "experimental": self.sources["experimental"]},
                        {"fe": self.sources["fe"], "experimental": None}):
            with self.subTest(sources=sources), self.assertRaisesRegex(rx.ExportRefused, "not recorded"):
                self.export(outcome, sources)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_a_changed_file_is_never_given_the_hash_of_the_read(self):
        # the report states the identity recorded at the read; a file that differs now is refused, not re-hashed
        outcome = self.run_compare()
        with self.changed(self.odb, lambda data: data[:-1] + b"#"):  # same size, other content
            with self.assertRaises(rx.ExportRefused):
                self.export(outcome)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_the_diagnostic_report_has_no_mac_and_no_pair(self):
        diagnostic = dm.run_diagnostic(self.fe_preview, self.exp_preview, 1, 1, threading.Event()).diagnostic
        written = rx.export_diagnostic(self.fe_preview, self.exp_preview, 1, diagnostic, "refused: test",
                                       self.folder, sources=self.sources)
        self.assertEqual(sorted(p.suffix for p in written), [".pdf", ".xlsx"])
        from openpyxl import load_workbook

        texts = " ".join(str(cell.value) for sheet in load_workbook(next(p for p in written if p.suffix == ".xlsx"))
                         for row in sheet.iter_rows() for cell in row if cell.value is not None)
        self.assertIn("NOT_COMPUTED", texts)
        self.assertNotIn("Excellent match", texts)
        with self.assertRaises(rx.ExportRefused):
            rx.export_diagnostic(self.fe_preview, self.exp_preview, 1, None, "", self.folder, sources=self.sources)

    def test_the_diagnostic_report_of_changed_or_unconfirmed_files_is_refused(self):
        diagnostic = dm.run_diagnostic(self.fe_preview, self.exp_preview, 1, 1, threading.Event()).diagnostic
        with self.changed_unv(lambda data: data + b"x"), self.assertRaises(rx.ExportRefused):
            rx.export_diagnostic(self.fe_preview, self.exp_preview, 1, diagnostic, "", self.folder,
                                 sources=self.sources)
        with self.assertRaises(rx.ExportRefused):
            rx.export_diagnostic(self.fe_preview, self.exp_preview, 1, diagnostic, "", self.folder, sources=None)
        self.assertEqual(list(self.folder.iterdir()), [])

    def test_a_missing_folder_is_refused(self):
        outcome = self.run_compare()
        with self.assertRaises(rx.ExportRefused):
            rx.export_verified(outcome.result, outcome.bound, outcome.view, self.fe_preview, self.exp_preview, 0,
                               self.folder / "missing", sources=outcome.sources)


class MaterialsTests(unittest.TestCase):
    def setUp(self):
        self.view = mm.load_material_evidence()

    def test_the_classes_are_separate_and_no_material_value_is_identified(self):
        self.assertEqual(self.view.problems, ())
        self.assertEqual(self.view.formal_output, mm.NO_GLOBAL_VALUE)
        identification = self.view.of(mm.MATERIAL_IDENTIFICATION)
        self.assertTrue(identification)
        self.assertEqual({row.value for row in identification}, {mm.NO_VALUE})
        self.assertTrue(all("NO_GLOBAL_PARAMETER_VALUE" in row.status for row in identification))
        self.assertEqual([row.status for row in self.view.of(mm.SPECIMEN_CALIBRATION)], [mm.NOT_AVAILABLE])
        self.assertEqual({row.status for row in self.view.of(mm.MODEL_INPUT)}, {"MODEL_INPUT"})
        for row in self.view.of(mm.DIAGNOSTIC_ESTIMATE):
            self.assertTrue("DIAGNOSTIC_ONLY" in row.status or "NOT_IDENTIFIABLE" in row.status
                            or "NOT_EVALUABLE" in row.status, row)
        candidate = next(r for r in self.view.of(mm.DIAGNOSTIC_ESTIMATE) if "OPTIMIZER_CANDIDATE" in r.role)
        self.assertIn("identified material property", candidate.reason)

    def test_the_model_inputs_are_the_governed_fe_models_constants(self):
        from domain.forward_model_manifest import load_forward_model_manifest

        constants = load_forward_model_manifest(
            ROOT / "docs/auto_id/forward_models/SP02.physical.forward.json").source_engineering_constants.to_dict()
        rows = {row.parameter: row for row in self.view.of(mm.MODEL_INPUT) if row.bound_to.startswith("SP02")}
        self.assertEqual({name: float(row.value) for name, row in rows.items()}, constants)
        self.assertEqual(rows["E1"].unit, "MPa")
        self.assertEqual(rows["nu12"].unit, "—")

    def test_no_record_is_attributed_to_the_read_data(self):
        # NEW-GUI.6-F1: a verified registration proves the geometry, source and set, not the FE model's identity
        self.assertEqual(mm.current_relevance(self.view, None, False).status, mm.NOT_AVAILABLE)
        rotated = mm.current_relevance(self.view, None, True)  # e.g. the SP-13 ROTATED series: no verified fixture
        self.assertEqual(rotated.status, mm.NOT_AVAILABLE)
        self.assertIn("not transferred", rotated.text)
        self.assertEqual(rotated.historical, "")
        for fixture in ("SP02/bravo-1-physical", "SP13/best-physical", "SP99/other"):
            with self.subTest(fixture=fixture):
                relevance = mm.current_relevance(self.view, fixture, True)
                self.assertEqual(relevance.status, mm.NOT_AVAILABLE)
                self.assertIn("No material record is attributed to the read data", relevance.text)
        sp02 = mm.current_relevance(self.view, "SP02/bravo-1-physical", True)
        self.assertIn("not a result of the read FE model", sp02.historical)
        self.assertIn("SP02_Modal_V02.inp SHA-256 574ae78a897f", sp02.historical)  # the record's own FE input
        self.assertIn(mm.NO_GLOBAL_VALUE, sp02.historical)  # D-077 shown unchanged, as history

    def test_a_record_that_does_not_match_its_pinned_hash_is_not_shown(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            shutil.copytree(ROOT / "docs" / "auto_id", root / "docs" / "auto_id")
            path = root / "docs/auto_id/campaigns/M7_RUN_A.result.json"
            record = json.loads(path.read_text(encoding="utf-8"))
            record["parameters"]["E_in_plane_mpa"] = 99999.0
            path.write_text(json.dumps(record), encoding="utf-8")
            view = mm.load_material_evidence(root)
        self.assertTrue(any("NOT VERIFIED" in text for text in view.problems))
        self.assertEqual(view.of(mm.MATERIAL_IDENTIFICATION), ())
        self.assertFalse(any("99999" in row.value for row in view.rows))

    def test_reading_writes_nothing(self):
        before = fingerprint(ROOT / "docs" / "auto_id")
        mm.load_material_evidence()
        self.assertEqual(fingerprint(ROOT / "docs" / "auto_id"), before)


class MaterialBindingTests(_Comparison):
    def test_another_fe_model_with_the_same_geometry_gets_no_material_binding(self):
        # Other frequencies (another material) on the same mesh: the registration applies, the M7 record does not.
        other = with_frequencies(self.fe, {m.number: 1.2 * m.frequency_hz for m in self.fe.sorted_modes()})
        fe_prev, exp_prev = self.previews(fe=other)
        check = rm.run_check(fe_prev, exp_prev, 0, 1, threading.Event())
        self.assertIsNotNone(check.view)  # the applicability check passes on the geometry
        self.assertEqual(check.view.fixture_id, "SP02/bravo-1-physical")
        relevance = mm.current_relevance(mm.load_material_evidence(), check.view.fixture_id, True)
        self.assertEqual(relevance.status, mm.NOT_AVAILABLE)
        self.assertNotIn(mm.NO_GLOBAL_VALUE, relevance.text)  # nothing of M7 is stated for the read FE model
        self.assertIn("not a result of the read FE model", relevance.historical)


class ReadIdentityTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.unv = F.write_frf_only(self.root / "frf.uff")

    def test_the_read_records_the_identity_of_the_file_read(self):
        from new_gui import workspace_state as ws

        outcome = pm.run_read(ws.EXPERIMENTAL, self.unv, 1, threading.Event())
        self.assertEqual(outcome.identity, pm.file_identity(self.unv))
        self.assertEqual(outcome.identity["size"], self.unv.stat().st_size)

    def test_a_file_changed_during_the_read_gets_no_identity(self):
        from new_gui import workspace_state as ws

        real = pm.read_preview

        def read_and_change(kind, path, cancel_event=None):
            preview = real(kind, path, cancel_event)
            with open(path, "ab") as stream:  # another program writes while the file is read
                stream.write(b"\n")
            return preview

        with mock.patch.object(pm, "read_preview", side_effect=read_and_change):
            outcome = pm.run_read(ws.EXPERIMENTAL, self.unv, 1, threading.Event())
        self.assertIsNotNone(outcome.preview)
        self.assertIsNone(outcome.identity)


class KnownSourceTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)

    def write_csv(self, change=None) -> Path:
        with open(ks.PROVENANCE, newline="", encoding="utf-8") as stream:
            rows = list(csv.DictReader(stream))
        if change:
            change(rows)
        path = self.root / "SOURCE_PROVENANCE.csv"
        with open(path, "w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        return path

    def test_the_accepted_map_confirms_exactly_the_three_fitted_sources(self):
        sources = ks.confirmed_sources()
        self.assertEqual([s.specimen for s in sources], ["SP-13", "SP-10", "PLATE-300x300x1mm"])
        self.assertTrue(all(s.relative_path.endswith(".unv") and "Polytec" in s.relative_path for s in sources))

    def test_a_status_token_in_a_description_is_not_a_working_source(self):
        def demote(rows):
            row = next(r for r in rows if r["sha256"] == ks.KNOWN_SOURCES[0].sha256)
            row["source_status"] = "AUXILIARY (formerly called NEW_WORKING_SOURCE in a draft)"

        with self.assertRaises(ks.ProvenanceMismatch):
            ks.confirmed_sources(self.write_csv(demote))

        def promote(rows):  # an FRF file whose description mentions the status is not added
            row = next(r for r in rows if "AUXILIARY_FRF" in r["source_status"])
            row["source_status"] = row["source_status"] + "; not a NEW_WORKING_SOURCE"

        self.assertEqual(len(ks.confirmed_sources(self.write_csv(promote))), 3)

    def test_another_sha_series_or_modal_set_type_is_refused(self):
        for field, value in (("sha256", "0" * 64), ("series_id", "NEXP29/SP-13/OTHER"),
                             ("file_type", "UNV: datasets [58, 151, 164, 2411]")):
            def change(rows, field=field, value=value):
                next(r for r in rows if r["sha256"] == ks.KNOWN_SOURCES[0].sha256)[field] = value

            with self.subTest(field=field), self.assertRaises(ks.ProvenanceMismatch):
                ks.confirmed_sources(self.write_csv(change))

    def test_a_missing_folder_or_a_changed_file_is_never_used(self):
        source = ks.KNOWN_SOURCES[0]
        self.assertEqual(ks.availability(source, self.root / "absent").state, ks.MISSING)
        with self.assertRaises(ks.ProvenanceMismatch):
            ks.verify(source, self.root / "absent")
        path = self.root / Path(source.relative_path.replace("\\", "/"))
        path.parent.mkdir(parents=True)
        path.write_bytes(b"\0" * source.size_bytes)  # the declared size, another content
        self.assertEqual(ks.availability(source, self.root).state, ks.AVAILABLE)
        with self.assertRaises(ks.ProvenanceMismatch):
            ks.verify(source, self.root)

    @unittest.skipUnless(ks.DEFAULT_ROOT.is_dir(), "the new experimental folder is not on this computer")
    def test_the_local_sources_are_verified_read_only(self):
        for source in ks.KNOWN_SOURCES:
            with self.subTest(source=source.specimen):
                path = ks.DEFAULT_ROOT / Path(source.relative_path.replace("\\", "/"))
                before = path.stat().st_mtime_ns
                self.assertEqual(ks.verify(source), path)
                self.assertEqual(path.stat().st_mtime_ns, before)


@unittest.skipUnless(REAL and Path(REAL).is_dir(), "NEW_GUI6_REAL_INPUTS (verified archived copies) is not set")
class RealArchivedTests(unittest.TestCase):
    """The positive path on the real archived SP-02 (bravo-1) and SP-13 (best) data."""

    CASES = (("SP02", "SP02_polymax_retry_260803.unv", "SP02/bravo-1-physical"),
             ("SP13", "SP13_a_polymax.unv", "SP13/best-physical"))

    def test_real_archived_verified_comparisons(self):
        from new_gui import preview_model as pm
        from new_gui import workspace_state as ws

        for specimen, unv, fixture in self.CASES:
            with self.subTest(specimen=specimen):
                root = Path(REAL) / specimen
                fe = pm.run_read(ws.FE, root / "fe_extraction" / "manifest.json", 1, threading.Event()).preview
                exp = pm.run_read(ws.EXPERIMENTAL, root / unv, 1, threading.Event()).preview
                index = next(i for i in range(len(exp.sets))
                             if rm.run_check(fe, exp, i, 1, threading.Event()).view is not None)
                check = rm.run_check(fe, exp, index, 1, threading.Event())
                self.assertEqual(check.view.fixture_id, fixture)
                outcome = cm.run_compare(fe, exp, index, check.view.registration, 1, threading.Event())
                self.assertEqual(outcome.error, "")
                direct = modal_core.compare_modal_datasets(fe.dataset, exp.sets[index].dataset,
                                                           frozen_registration=outcome.bound.registration)
                self.assertEqual([(p.abaqus_mode, p.experimental_mode, p.mac) for p in outcome.result.pairs],
                                 [(p.abaqus_mode, p.experimental_mode, p.mac) for p in direct.pairs])
                self.assertTrue(outcome.result.pairs)
                self.assertEqual(outcome.view.status, cm.PAIRS_FOUND)
                images = cm.run_pair_images(outcome.result, 0, 1, threading.Event())
                self.assertEqual(len(images.images), 3)
                other = [i for i in range(len(exp.sets)) if i != index]
                for i in other:  # another set of the same file is refused
                    self.assertIsNone(cm.run_compare(fe, exp, i, check.view.registration, 1,
                                                     threading.Event()).view)


class StartTests(unittest.TestCase):
    def test_importing_the_models_loads_no_comparator_or_report_writer(self):
        import subprocess

        code = ("import sys; sys.path.insert(0, sys.argv[1]); "
                "import new_gui.comparison_model, new_gui.materials_model, new_gui.known_sources, "
                "new_gui.report_export; "
                "bad = sorted(m for m in sys.modules if m in ('reviewed_core', 'quality_control', 'reporting', "
                "'amplitude_correlation', 'reporting_hardening', 'main', 'tkinter')); print(','.join(bad))")
        done = subprocess.run([sys.executable, "-c", code, str(ROOT / "src")], capture_output=True, text=True,
                              timeout=120, cwd=str(ROOT), env={**os.environ, "PYTHONIOENCODING": "utf-8",
                                                               "PYTHONDONTWRITEBYTECODE": "1"})
        self.assertEqual(done.returncode, 0, done.stderr)
        self.assertEqual(done.stdout.strip(), "")


if __name__ == "__main__":
    unittest.main()
