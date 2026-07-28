from pathlib import Path
import sys
import tempfile
import unittest

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import main  # noqa: F401  # installs the reviewed runtime/reporting stack
from modal_core import ModalDataset, ModeShape
from quality_control_reviewed import compare_modal_datasets_with_quality_control
from reporting import export_excel, export_pdf


class ReviewedIntegrationTests(unittest.TestCase):
    def test_reviewed_stack_imports_and_exports_reports(self):
        coordinates = np.array(
            [
                [0.0, 0.0, 0.0],
                [2.0, 0.0, 0.0],
                [0.0, 1.0, 0.0],
                [2.0, 1.0, 0.0],
                [0.6, 0.3, 0.0],
                [1.4, 0.7, 0.0],
            ]
        )
        node_ids = np.arange(1, len(coordinates) + 1)
        z1 = np.array([1.0, -1.0, -1.0, 1.0, 0.4, -0.4])
        z2 = np.array([1.0, 1.0, -1.0, -1.0, 0.2, -0.2])

        abaqus_modes = []
        experimental_modes = []
        for number, (frequency_a, frequency_e, shape) in enumerate(
            ((50.5, 50.0, z1), (101.0, 100.0, z2)), start=1
        ):
            abaqus_vectors = np.column_stack((10.0 * shape, 5.0 * shape, shape))
            experimental_vectors = np.zeros_like(abaqus_vectors)
            experimental_vectors[:, 2] = shape
            abaqus_modes.append(
                ModeShape(number, frequency_a, node_ids, coordinates, abaqus_vectors)
            )
            experimental_mode = ModeShape(
                number, frequency_e, node_ids, coordinates, experimental_vectors
            )
            experimental_mode.measured_dofs = np.zeros_like(
                experimental_vectors, dtype=bool
            )
            experimental_mode.measured_dofs[:, 2] = True
            experimental_modes.append(experimental_mode)

        result = compare_modal_datasets_with_quality_control(
            ModalDataset("Abaqus", Path("model.odb"), abaqus_modes),
            ModalDataset("Experiment", Path("scan.unv"), experimental_modes),
            coordinate_scale_override=1.0,
        )
        self.assertEqual(len(result.pairs), 2)
        self.assertTrue(all(pair.mac > 0.999 for pair in result.pairs))
        self.assertTrue(all(pair.frequency_error_percent > 0.0 for pair in result.pairs))

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            excel = export_excel(result, root / "report.xlsx", root / "images")
            pdf = export_pdf(result, root / "report.pdf", root / "images")
            self.assertGreater(excel.stat().st_size, 1000)
            self.assertGreater(pdf.stat().st_size, 1000)

            workbook = load_workbook(excel, read_only=True)
            try:
                self.assertIn("AutoMAC COMAC", workbook.sheetnames)
                self.assertIn("FRF Diagnostics", workbook.sheetnames)
                self.assertEqual(
                    workbook["Mode Comparison"].cell(row=1, column=5).value,
                    "Signed frequency error, %",
                )
                summary = workbook["Summary"]
                # Mean MAC (written by the base reporting.export_excel) must not be
                # clobbered by the frequency-error summary rows the reporting
                # extension layers add afterward.
                self.assertEqual(summary.cell(row=12, column=1).value, "Mean MAC")
                self.assertAlmostEqual(
                    summary.cell(row=12, column=2).value,
                    float(np.mean([pair.mac for pair in result.pairs])),
                    places=6,
                )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
